"""PPTX export of plot/figure objects (docs/FIGURES_CONTRACT.md 4x): plain pictures, placeholders, warnings."""
import base64
import io
import json

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.exporter import EMU_PER_PIXEL, export_presentation, to_emu
from backend.models import PresentationDoc, SlideModel, SlideObject


def png_bytes(w, h, seed=0):
    img = Image.new("RGB", (w, h))
    px = img.load()
    for y in range(h):
        for x in range(w):
            px[x, y] = ((x * 3 + seed) % 256, (y * 5) % 256, ((x + y) * 2) % 256)
    buf = io.BytesIO()
    img.save(buf, "PNG", dpi=(450, 450))
    return buf.getvalue()


def data_url(raw):
    return "data:image/png;base64," + base64.b64encode(raw).decode()


def export(objs):
    warnings = []
    doc = PresentationDoc(title="Fig deck", themeId="slate", slides=[SlideModel(id="s1", layout="blank", objects=objs)])
    prs = Presentation(export_presentation(doc, warnings=warnings))
    return prs.slides[0], warnings


def obj(oid, typ, src, w, h, **kw):
    return SlideObject(id=oid, type=typ, figureId="fig_" + oid, src=src, x=50, y=60, w=w, h=h, **kw)


def test_plot_and_composite_export_as_plain_pictures():
    p1, p2 = png_bytes(640, 400), png_bytes(600, 600, seed=7)
    slide, warnings = export([
        obj("a", "plot", data_url(p1), 480, 300, name="Dose response", alt="Line plot of dose response."),
        obj("b", "figure", data_url(p2), 300, 300, name="Panel A-D"),
    ])
    assert warnings == []
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 2
    a, b = pics
    assert a.image.blob == p1 and b.image.blob == p2          # bytes unchanged, resolution intact
    assert (a.width, a.height) == (to_emu(480), to_emu(300))   # picture size == object geometry
    assert (a.left, a.top) == (to_emu(50), to_emu(60))
    assert (b.width, b.height) == (to_emu(300), to_emu(300))
    for pic in pics:                                            # no cropping
        assert (pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom) == (0, 0, 0, 0)
    assert a.name == "Figure \u2013 Dose response" and b.name == "Figure \u2013 Panel A-D"
    assert a._element.nvPicPr.cNvPr.get("descr") == "Line plot of dose response."
    assert b._element.nvPicPr.cNvPr.get("descr") == ""         # alt only from object.alt, never a file name
    assert "figureId" not in a._element.xml and "fig_a" not in a._element.xml  # plain image on re-import


def test_mismatched_frame_is_never_stretched():
    raw = png_bytes(640, 400)
    slide, warnings = export([obj("a", "plot", data_url(raw), 400, 400, name="Skewed")])
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert pic.width == to_emu(400) and abs(pic.height - pic.width * 400 / 640) <= 1
    assert [w["code"] for w in warnings] == ["figure_aspect_adjusted"]
    # sub-pixel rounding of the stored height is not a mismatch
    slide, warnings = export([obj("a", "plot", data_url(raw), 333, 208, name="Rounded")])
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert warnings == [] and pic.height == to_emu(208)


@pytest.mark.parametrize("src,code", [(None, "figure_image_missing"), ("", "figure_image_missing"),
                                      ("data:image/png;base64,AAAA", "figure_image_unreadable"),
                                      ("https://example.com/x.png", "figure_image_unreadable")])
def test_missing_src_gives_visible_placeholder_and_warning(src, code):
    slide, warnings = export([obj("m", "plot", src, 400, 250, name="Growth curve", alt="Growth curve.")])
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    boxes = [s for s in slide.shapes if s.has_text_frame and "Growth curve" in s.text_frame.text]
    assert len(boxes) == 1
    box = boxes[0]
    assert (box.left, box.top, box.width, box.height) == (to_emu(50), to_emu(60), to_emu(400), to_emu(250))
    assert box.text_frame.text.strip().splitlines()[0] == "Growth curve" and len(box.text_frame.text) > len("Growth curve")
    assert len(warnings) == 1 and warnings[0]["code"] == code
    assert warnings[0]["slide"] == 1 and warnings[0]["objectId"] == "m" and "Growth curve" in warnings[0]["message"]


def test_placeholder_uses_id_when_unnamed():
    slide, warnings = export([obj("m", "figure", None, 200, 100)])
    assert warnings[0]["name"] == "fig_m"
    assert any("fig_m" in s.text_frame.text for s in slide.shapes if s.has_text_frame)


def test_other_objects_still_export_next_to_figures():
    slide, warnings = export([
        SlideObject(id="t", type="text", x=10, y=10, w=200, h=40, text="Hello"),
        obj("a", "plot", data_url(png_bytes(64, 40)), 128, 80, name="Tiny"),
    ])
    assert warnings == [] and EMU_PER_PIXEL == 9525
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    texts = [s for s in slide.shapes if s.has_text_frame and "Hello" in s.text_frame.text]
    assert len(pics) == 1 and len(texts) == 1


def test_endpoint_header_lists_warnings():
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    from backend.server import app

    client = TestClient(app)

    def post(objs):
        doc = PresentationDoc(title="T", themeId="slate", slides=[SlideModel(id="s1", layout="blank", objects=objs)])
        return client.post("/api/presentation/export", json={"format": "pptx", "doc": doc.model_dump()})

    res = post([obj("m", "plot", None, 400, 250, name="Gr\u00f6\u00dfe")])
    assert res.status_code == 200
    payload = json.loads(res.headers["x-export-warnings"])
    assert payload["count"] == 1 and payload["items"][0]["code"] == "figure_image_missing"
    assert payload["items"][0]["name"] == "Gr\u00f6\u00dfe"      # non-ASCII survives the ASCII-only header
    assert "X-Export-Warnings" in res.headers["access-control-expose-headers"]
    ok = post([obj("a", "plot", data_url(png_bytes(64, 40)), 128, 80, name="Fine")])
    assert ok.status_code == 200 and "x-export-warnings" not in ok.headers
