"""Regression tests: list-paragraph spacing in PPTX export.

Root-cause fixed: The slide-master bodyStyle defines spcBef=20% for every bullet
level.  Without an explicit <a:spcBef> override in <a:pPr>, PowerPoint inherits
that 20 % of the font size as blank space before each list item – producing the
visual "blank line" between bullets even though the browser preview is fine.

Fix: consecutive list paragraphs now emit
  <a:spcBef><a:spcPts val="0"/></a:spcBef>
  <a:spcAft><a:spcPts val="0"/></a:spcAft>
inside their pPr, suppressing inherited spacing.  Intentional blank paragraphs,
prose paragraphs and explicit obj.spaceAfter values are preserved.
"""
import io

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from backend.exporter import export_presentation
from backend.models import PresentationDoc, SlideModel, SlideObject

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_doc(obj: SlideObject) -> PresentationDoc:
    return PresentationDoc(
        title="Spacing test",
        themeId="journal",
        slides=[SlideModel(id="s1", layout="blank", objects=[obj])],
    )


def _export(obj: SlideObject) -> Presentation:
    return Presentation(export_presentation(_make_doc(obj)))


def _paras(prs: Presentation):
    """Return all a:p elements from the last shape on slide 0."""
    return prs.slides[0].shapes[-1].text_frame.paragraphs


def _spc_pts(p, tag: str) -> int | None:
    """Return the integer spcPts value for spcBef or spcAft on a paragraph, or None."""
    ppr = p._p.find(f"{{{A_NS}}}pPr")
    if ppr is None:
        return None
    el = ppr.find(f"{{{A_NS}}}{tag}")
    if el is None:
        return None
    child = el.find(f"{{{A_NS}}}spcPts")
    if child is None:
        return None
    return int(child.get("val", "0"))


def _has_bullet_char(p) -> bool:
    ppr = p._p.find(f"{{{A_NS}}}pPr")
    return ppr is not None and ppr.find(f"{{{A_NS}}}buChar") is not None


def _has_bu_auto_num(p) -> bool:
    ppr = p._p.find(f"{{{A_NS}}}pPr")
    return ppr is not None and ppr.find(f"{{{A_NS}}}buAutoNum") is not None


def _has_bu_none(p) -> bool:
    ppr = p._p.find(f"{{{A_NS}}}pPr")
    return ppr is not None and ppr.find(f"{{{A_NS}}}buNone") is not None


# ---------------------------------------------------------------------------
# 1. Consecutive bullet items – no inherited spcBef / spcAft
# ---------------------------------------------------------------------------

def test_bullet_list_spc_bef_is_zero():
    """Every bullet-list paragraph must explicitly zero spcBef to suppress the
    bodyStyle 20% inherited from the slide master."""
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=18.0,
        html="<ul><li>Alpha</li><li>Bravo</li><li>Charlie</li></ul>",
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    assert len(paras) == 3
    for p in paras:
        assert _has_bullet_char(p), f"Expected bullet char on: {p.text!r}"
        assert _spc_pts(p, "spcBef") == 0, (
            f"spcBef should be 0 for list item {p.text!r}; "
            "inherited bodyStyle 20% would cause blank-line gaps"
        )
        assert _spc_pts(p, "spcAft") == 0, f"spcAft should be 0 for list item {p.text!r}"


def test_numbered_list_spc_bef_is_zero():
    """Numbered list items must also suppress inherited spcBef."""
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=18.0,
        html="<ol><li>Step one</li><li>Step two</li><li>Step three</li></ol>",
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    assert len(paras) == 3
    for p in paras:
        assert _has_bu_auto_num(p), f"Expected auto-num on: {p.text!r}"
        assert _spc_pts(p, "spcBef") == 0, f"spcBef not zeroed for numbered item {p.text!r}"
        assert _spc_pts(p, "spcAft") == 0, f"spcAft not zeroed for numbered item {p.text!r}"


# ---------------------------------------------------------------------------
# 2. Mixed list types in a single textbox
# ---------------------------------------------------------------------------

def test_mixed_bullet_and_numbered_list():
    """Bullet items followed immediately by numbered items must all have spcBef=0."""
    html = "<ul><li>Bullet A</li><li>Bullet B</li></ul><ol><li>Num 1</li><li>Num 2</li></ol>"
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=16.0,
        html=html,
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    assert len(paras) == 4
    # First two: bullets
    for p in paras[:2]:
        assert _has_bullet_char(p)
        assert _spc_pts(p, "spcBef") == 0
    # Last two: numbers
    for p in paras[2:]:
        assert _has_bu_auto_num(p)
        assert _spc_pts(p, "spcBef") == 0


# ---------------------------------------------------------------------------
# 3. Nested list (ul inside li) — children also zeroed
# ---------------------------------------------------------------------------

def test_nested_list_spc_bef_is_zero():
    """Nested ul inside a li produces multiple bullet paragraphs; all must have spcBef=0."""
    html = (
        "<ul>"
        "<li>Parent A<ul><li>Child A1</li><li>Child A2</li></ul></li>"
        "<li>Parent B</li>"
        "</ul>"
    )
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=400,
        fontSize=16.0,
        html=html,
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    # All resulting paragraphs are list paragraphs; none should inherit master spcBef.
    for p in paras:
        if _has_bullet_char(p) or _has_bu_auto_num(p):
            assert _spc_pts(p, "spcBef") == 0, f"Nested item {p.text!r} inherits spcBef"


# ---------------------------------------------------------------------------
# 4. Intentional blank paragraph in a list is preserved
# ---------------------------------------------------------------------------

def test_blank_paragraph_between_list_items_preserved():
    """A paragraph with no runs is an intentional separator and must be kept."""
    # A <p></p> between two <li>s produces an empty paragraph in the data model.
    html = "<ul><li>Item 1</li></ul><p></p><ul><li>Item 2</li></ul>"
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=16.0,
        html=html,
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    # We should have: item1, blank, item2
    assert len(paras) >= 3, "Blank paragraph between list groups must be preserved"
    texts = [p.text for p in paras]
    assert "Item 1" in texts
    assert "Item 2" in texts


# ---------------------------------------------------------------------------
# 5. Prose paragraphs are NOT affected
# ---------------------------------------------------------------------------

def test_prose_paragraphs_not_affected():
    """Plain text paragraphs must NOT have spcBef forced to zero – they should
    inherit theme spacing or reflect obj.spaceAfter as intended."""
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=200,
        fontSize=16.0,
        html="<p>First prose paragraph.</p><p>Second prose paragraph.</p>",
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    # Prose paragraphs must have buNone (not bullet markers) and must NOT have
    # an explicit spcBef=0 (that would be the list-item suppression leaking).
    for p in paras:
        assert _has_bu_none(p), f"Prose paragraph {p.text!r} should have buNone"
        # For prose, spcBef is either None (not written) or inherited; we must
        # NOT have written an explicit zero that would clobber theme spacing.
        spc = _spc_pts(p, "spcBef")
        assert spc is None, (
            f"Prose paragraph {p.text!r} has an unexpected explicit spcBef={spc}; "
            "the list-item zero must not bleed onto prose paragraphs"
        )


# ---------------------------------------------------------------------------
# 6. Explicit obj.spaceAfter is honoured for list items
# ---------------------------------------------------------------------------

def test_explicit_space_after_honoured_for_list_items():
    """If obj.spaceAfter is set, list items should emit that value in spcAft,
    not zero, so the author's vertical rhythm is respected."""
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=18.0,
        html="<ul><li>Item X</li><li>Item Y</li></ul>",
        lineHeight=1.2, spaceAfter=6.0,  # 6pt gap after each list item
    )
    prs = _export(obj)
    paras = _paras(prs)
    for p in paras:
        if _has_bullet_char(p):
            aft = _spc_pts(p, "spcAft")
            assert aft == 600, (  # 6.0pt × 100 = 600 spcPts units
                f"spcAft should be 600 (6pt) for list item with spaceAfter=6; got {aft}"
            )
            bef = _spc_pts(p, "spcBef")
            assert bef == 0, f"spcBef should still be 0 regardless of spaceAfter; got {bef}"


# ---------------------------------------------------------------------------
# 7. Mixed formatting inside a list item is preserved
# ---------------------------------------------------------------------------

def test_mixed_formatting_in_list_item():
    """Bold/italic inline formatting inside list items must survive the fix."""
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=16.0,
        html="<ul><li><b>Bold</b> and <i>italic</i> text</li><li>Normal item</li></ul>",
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    assert len(paras) == 2
    # First para must have multiple runs (bold + normal + italic)
    runs0 = paras[0]._p.findall(f"{{{A_NS}}}r")
    assert len(runs0) >= 2, "Bold/italic runs in list item should be preserved"
    # spcBef still 0
    assert _spc_pts(paras[0], "spcBef") == 0
    assert _spc_pts(paras[1], "spcBef") == 0


# ---------------------------------------------------------------------------
# 8. obj.list="bullet" (non-HTML bullet mode) also zeroed
# ---------------------------------------------------------------------------

def test_obj_level_bullet_mode():
    """Textbox with obj.list='bullet' (plain-text list) must also zero spcBef."""
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=200,
        fontSize=16.0,
        text="Line one\nLine two\nLine three",
        list="bullet",
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    for p in paras:
        if p.text and p.text.strip():
            assert _has_bullet_char(p), f"Expected bullet char on: {p.text!r}"
            assert _spc_pts(p, "spcBef") == 0, f"spcBef not zeroed: {p.text!r}"


# ---------------------------------------------------------------------------
# 9. Bullet text carries correct runs / content
# ---------------------------------------------------------------------------

def test_bullet_run_text_is_correct():
    """Bullet export must not corrupt run text."""
    items = ["Alpha", "Bravo", "Charlie Delta"]
    html = "<ul>" + "".join(f"<li>{t}</li>" for t in items) + "</ul>"
    obj = SlideObject(
        id="t1", type="text", x=50, y=50, w=500, h=300,
        fontSize=18.0, html=html,
        lineHeight=1.2, spaceAfter=0.0,
    )
    prs = _export(obj)
    paras = _paras(prs)
    assert len(paras) == len(items)
    for p, expected in zip(paras, items):
        assert p.text == expected, f"Expected {expected!r}, got {p.text!r}"
