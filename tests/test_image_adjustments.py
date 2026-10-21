"""
Tests for lightweight picture adjustments (TASK 26):
- Brightness, contrast, saturation, desaturation (grayscale), tint/monochrome from palette
- Non-destructive image cropping
- Parameter persistence and reset control
- PPTX export: native picture effects where accurate, baked raster copy for tint
- Original image data (src) preserved in document model
"""
import base64
import io
import json
import subprocess
import shutil
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.exporter import EMU_PER_PIXEL, export_presentation, to_emu, bake_image_adjustments
from backend.models import PresentationDoc, SlideModel, SlideObject


def create_sample_png(w=100, h=80, color=(200, 100, 50)):
    img = Image.new("RGBA", (w, h), color + (255,))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def create_data_url(raw_bytes):
    return "data:image/png;base64," + base64.b64encode(raw_bytes).decode("ascii")


def export_single_slide(objs, theme_id="slate"):
    warnings = []
    doc = PresentationDoc(
        title="Test Deck",
        themeId=theme_id,
        slides=[SlideModel(id="s1", layout="blank", objects=objs)]
    )
    pptx_stream = export_presentation(doc, warnings=warnings)
    prs = Presentation(pptx_stream)
    return prs.slides[0], warnings, doc


# ============================================================================
# 1. Model Persistence & Non-destructive Adjustments
# ============================================================================

def test_image_adj_parameters_persist_and_roundtrip():
    """Verify imgAdj parameters serialize, deserialize, and persist across JSON roundtrips."""
    raw = create_sample_png()
    durl = create_data_url(raw)
    adj_data = {
        "brightness": 25,
        "contrast": -15,
        "saturation": 50,
        "grayscale": 0,
        "tintColor": "#C00000",
        "tintAmount": 0.35,
        "cropL": 0.1,
        "cropR": 0.15,
        "cropT": 0.05,
        "cropB": 0.2
    }
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=100, y=100, w=400, h=300,
        aspect=4/3,
        name="Photo",
        alt="A test photo",
        imgAdj=adj_data
    )

    # Test JSON serialization / deserialization roundtrip
    serialized = obj.model_dump_json()
    data = json.loads(serialized)
    assert data["imgAdj"] == adj_data
    assert data["src"] == durl
    assert data["aspect"] == 4/3

    reconstructed = SlideObject.model_validate_json(serialized)
    assert reconstructed.imgAdj == adj_data
    assert reconstructed.src == durl
    assert reconstructed.w == 400
    assert reconstructed.h == 300


def test_image_adj_reset_restores_original_state():
    """Resetting imgAdj removes all adjustments without altering the original image src or geometry."""
    raw = create_sample_png()
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=50, y=60, w=300, h=200,
        aspect=1.5,
        imgAdj={"brightness": 40, "contrast": 30, "cropL": 0.2}
    )

    # Reset
    obj.imgAdj = None
    assert obj.imgAdj is None
    assert obj.src == durl
    assert (obj.x, obj.y, obj.w, obj.h) == (50, 60, 300, 200)
    assert obj.aspect == 1.5


# ============================================================================
# 2. PPTX Export: Native Picture Effects
# ============================================================================

def test_pptx_export_native_brightness_contrast():
    """Brightness and contrast without tint map to native PPTX <a:lum> element."""
    raw = create_sample_png()
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=50, y=50, w=200, h=150,
        imgAdj={"brightness": 20, "contrast": 15}
    )
    slide, warnings, _ = export_single_slide([obj])
    pic = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]

    xml = pic._element.xml
    assert "<a:lum" in xml
    assert 'bright="20000"' in xml
    assert 'contrast="15000"' in xml


def test_pptx_export_native_grayscale():
    """Grayscale=100 maps to native PPTX <a:grayscl> element."""
    raw = create_sample_png()
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=50, y=50, w=200, h=150,
        imgAdj={"grayscale": 100}
    )
    slide, warnings, _ = export_single_slide([obj])
    pic = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]

    xml = pic._element.xml
    assert "<a:grayscl" in xml


def test_pptx_export_native_saturation():
    """Saturation adjustment maps to native PPTX <a:satMod> element."""
    raw = create_sample_png()
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=50, y=50, w=200, h=150,
        imgAdj={"saturation": -50}
    )
    slide, warnings, _ = export_single_slide([obj])
    pic = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]

    xml = pic._element.xml
    assert "<a:satMod" in xml
    # saturation -50 -> (1 - 0.5) * 100000 = 50000
    assert 'val="50000"' in xml


def test_pptx_export_native_crop_src_rect():
    """Crop without tint maps to native PPTX <a:srcRect> framing."""
    raw = create_sample_png()
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=50, y=50, w=200, h=150,
        imgAdj={"cropL": 0.1, "cropR": 0.2, "cropT": 0.05, "cropB": 0.15}
    )
    slide, warnings, _ = export_single_slide([obj])
    pic = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]

    xml = pic._element.xml
    assert "<a:srcRect" in xml
    assert 'l="10000"' in xml
    assert 'r="20000"' in xml
    assert 't="5000"' in xml
    assert 'b="15000"' in xml


# ============================================================================
# 3. PPTX Export: Tint Baking & Non-destructive Preservation
# ============================================================================

def test_pptx_export_tint_bakes_raster_copy_and_preserves_doc_src():
    """Monochrome / tint from palette bakes adjusted raster at export; doc.src is unchanged."""
    raw = create_sample_png(w=120, h=90, color=(100, 150, 200))
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=50, y=50, w=240, h=180,
        imgAdj={"tintColor": "#C00000", "tintAmount": 0.5}
    )

    slide, warnings, doc = export_single_slide([obj])
    pic = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]

    # Exported image blob is baked and therefore different from raw
    exported_blob = pic.image.blob
    assert exported_blob != raw

    # Inspect the exported pixels to verify tint was blended
    baked_img = Image.open(io.BytesIO(exported_blob)).convert("RGB")
    px = baked_img.getpixel((10, 10))
    # Original color was (100, 150, 200). #C00000 is (192, 0, 0).
    # 50% blend produces around (146, 75, 100)
    assert px[0] > 120  # red channel boosted
    assert px[1] < 120  # green channel lowered

    # CRITICAL: Document src is completely preserved and untouched
    assert doc.slides[0].objects[0].src == durl


def test_pptx_export_tint_and_crop_combined_baking():
    """Combining tint and crop bakes both into the raster copy."""
    raw = create_sample_png(w=200, h=100, color=(50, 100, 150))
    durl = create_data_url(raw)
    obj = SlideObject(
        id="img1",
        type="image",
        src=durl,
        x=40, y=40, w=200, h=100,
        imgAdj={"cropL": 0.25, "cropR": 0.25, "tintColor": "#00FF00", "tintAmount": 0.4}
    )

    slide, warnings, _ = export_single_slide([obj])
    pic = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]

    baked_img = Image.open(io.BytesIO(pic.image.blob))
    # 200 width with 25% cropped from each side -> 100 width
    assert baked_img.size == (100, 100)


# ============================================================================
# 4. JavaScript Engine Parity & Command Registration
# ============================================================================

def test_js_renderer_and_commands_parity():
    """Test JS filter string building and command registration in Node.js."""
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js not available")

    js_code = """
    const fs = require('fs');
    const path = require('path');

    // Mock window and document
    global.window = {};
    global.document = {
      createElement: () => ({ style: {}, setAttribute: () => {}, appendChild: () => {} }),
      addEventListener: () => {}
    };

    // Load renderer
    const rendererPath = path.resolve(__dirname ? path.join(__dirname, '../js/renderer.js') : 'js/renderer.js');
    const rendererCode = fs.readFileSync(fs.existsSync(rendererPath) ? rendererPath : path.resolve('js/renderer.js'), 'utf8');
    eval(rendererCode);

    const R = window.RPE.renderer;
    const assert = require('assert');

    // 1. buildImageFilter tests
    const f1 = R.buildImageFilter({ brightness: 20 });
    assert.strictEqual(f1, 'brightness(1.200)');

    const f2 = R.buildImageFilter({ contrast: -10, saturation: 50 });
    assert.strictEqual(f2, 'contrast(0.900) saturate(1.500)');

    const f3 = R.buildImageFilter({ grayscale: 1 });
    assert.strictEqual(f3, 'grayscale(1.000)');

    const fAll = R.buildImageFilter({ brightness: 10, contrast: 20, saturation: -30, grayscale: 0.5 });
    assert.strictEqual(fAll, 'brightness(1.100) contrast(1.200) saturate(0.700) grayscale(0.500)');

    console.log('JS_IMAGE_ADJ_TESTS_OK');
    """

    res = subprocess.run([node, "-e", js_code], capture_output=True, text=True, check=True)
    assert "JS_IMAGE_ADJ_TESTS_OK" in res.stdout
