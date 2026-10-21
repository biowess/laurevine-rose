"""End-to-end integration and visual regression tests for Scientific Figure Studio.

Covers:
- Full backend e2e pipeline for all 6 plot types (preview, final@450dpi, export png/pdf/py,
  standalone execution, pixel parity, slide placement, PPTX export, LibreOffice render check).
- Composite figure e2e for all 6 layouts.
- Python custom figure e2e for all starter templates.
- LaTeX figure e2e for all templates (skipped if TeX missing).
- Visual regression against downscaled grayscale perceptual fingerprint fixtures.
- Parity verification: generated standalone script vs endpoint render.
- Failure handling: malformed CSV, log-axis with zeros, empty data, timeout,
  missing module, and missing TeX compiler (all producing ok:false with human messages).
"""
from __future__ import annotations

import base64
import copy
import io
import os
import shutil
import subprocess
import sys
import tempfile
from typing import Any, Dict, List

import numpy as np
import pytest
from PIL import Image
from pptx import Presentation
from starlette.testclient import TestClient

from backend.figures import export as export_core, jobs, latex as latex_core, pyrun, render as render_core, runner, samples, templates
from backend.figures.models import default_record
from backend.figures.runtime import get_runtime
from backend.models import PresentationDoc, SlideModel, SlideObject
from backend.server import app
from tests.fixtures.figures.regenerate_fixtures import (
    compare_fingerprints,
    compute_fingerprint,
    load_fixtures,
    make_sample_record,
)

client = TestClient(app)

PLOT_TYPES = list(samples.PLOT_TYPES)  # ["line", "scatter", "bar", "hist", "box", "heatmap"]
COMPOSITE_LAYOUTS = ["2x2", "1x2", "2x1", "1x3", "1+2", "2x3"]


@pytest.fixture(autouse=True)
def _clean_figure_state():
    render_core.clear_cache()
    jobs.reset()
    yield
    jobs.reset()


def decode_png(preview_url: str) -> Image.Image:
    assert preview_url.startswith("data:image/png;base64,")
    raw = base64.b64decode(preview_url.split(",", 1)[1])
    img = Image.open(io.BytesIO(raw))
    img.load()
    return img


# =====================================================================
# 1. Visual Regression (Perceptual Fingerprints)
# =====================================================================

@pytest.mark.parametrize("plot_type", PLOT_TYPES)
def test_visual_regression_fingerprints(plot_type: str):
    """Ensure structured plots match downscaled grayscale fingerprints within tolerance."""
    fixtures = load_fixtures()
    ref_info = fixtures["plots"].get(plot_type)
    assert ref_info is not None, f"No reference fixture for {plot_type}"

    rec = make_sample_record(plot_type, w=4.0, h=3.0, dpi=100)
    res = client.post("/api/figures/render", json={"record": rec, "purpose": "preview", "dpi": 100})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is True and body["status"] == "ok"

    img = decode_png(body["preview"])
    actual_fp = compute_fingerprint(img)
    metrics = compare_fingerprints(actual_fp, ref_info["fingerprint"])

    # Grayscale pixel differences: allow tolerance for anti-aliasing variations (<= 12.0 out of 255)
    # A styling regression (wrong colors, missing axes, inverted scales) produces mean differences > 25.
    assert metrics["mean_diff"] <= 12.0, (
        f"Visual regression detected for '{plot_type}': "
        f"mean diff {metrics['mean_diff']:.2f} > 12.0 (rmse: {metrics['rmse']:.2f})"
    )


# =====================================================================
# 2. Backend E2E: 6 Plot Types Full Lifecycle
# =====================================================================

@pytest.mark.slow
@pytest.mark.parametrize("plot_type", PLOT_TYPES)
def test_plot_type_e2e_full_lifecycle(plot_type: str):
    """Full lifecycle e2e test for each of the 6 plot types:
    record -> /render preview -> /render final@450dpi -> /export png/pdf/py ->
    run .py standalone -> compare pixel size -> place in doc -> /presentation/export PPTX ->
    LibreOffice render check.
    """
    w_in, h_in = 4.0, 3.0
    rec = make_sample_record(plot_type, w=w_in, h=h_in, dpi=150)
    expected_w_px = round(w_in * 450)
    expected_h_px = round(h_in * 450)

    # 1. /render preview
    res_prev = client.post("/api/figures/render", json={"record": rec, "purpose": "preview"})
    assert res_prev.status_code == 200, res_prev.text
    body_prev = res_prev.json()
    assert body_prev["ok"] is True and body_prev["status"] == "ok"
    img_prev = decode_png(body_prev["preview"])
    # Preview DPI is capped at 150
    assert img_prev.width <= round(w_in * 150) + 1 and img_prev.height <= round(h_in * 150) + 1

    # 2. /render final@450dpi
    res_final = client.post("/api/figures/render", json={"record": rec, "purpose": "final", "dpi": 450})
    assert res_final.status_code == 200, res_final.text
    body_final = res_final.json()
    assert body_final["ok"] is True and body_final["status"] == "ok"
    assert (body_final["widthPx"], body_final["heightPx"]) == (expected_w_px, expected_h_px)
    final_data_url = body_final["preview"]

    # 3. /export png
    res_png = client.post("/api/figures/export", json={"record": rec, "format": "png", "dpi": 450})
    assert res_png.status_code == 200, res_png.text
    assert res_png.headers["content-type"] == "image/png"
    img_png = Image.open(io.BytesIO(res_png.content))
    img_png.load()
    assert img_png.size == (expected_w_px, expected_h_px)

    # 4. /export pdf
    res_pdf = client.post("/api/figures/export", json={"record": rec, "format": "pdf"})
    assert res_pdf.status_code == 200, res_pdf.text
    assert res_pdf.content.startswith(b"%PDF")
    assert res_pdf.headers["content-type"] == "application/pdf"

    # 5. /export py
    res_py = client.post("/api/figures/export", json={"record": rec, "format": "py", "dpi": 450})
    assert res_py.status_code == 200, res_py.text
    script_content = res_py.content.decode("utf-8")
    assert "import matplotlib.pyplot as plt" in script_content
    assert "DPI = 450" in script_content

    # 6. Run exported .py standalone & compare pixel size
    with tempfile.TemporaryDirectory() as tmpdir:
        py_path = os.path.join(tmpdir, "exported_figure.py")
        with open(py_path, "w", encoding="utf-8") as f:
            f.write(script_content)

        env = os.environ.copy()
        env["MPLBACKEND"] = "Agg"
        proc = subprocess.run(
            [sys.executable, py_path],
            cwd=tmpdir,
            env=env,
            capture_output=True,
            text=True,
            timeout=30,
        )
        assert proc.returncode == 0, f"Standalone .py failed:\nSTDOUT:\n{proc.stdout}\nSTDERR:\n{proc.stderr}"

        png_files = [f for f in os.listdir(tmpdir) if f.endswith(".png")]
        assert len(png_files) >= 1, f"No PNG generated by standalone script in {tmpdir}"
        standalone_png_path = os.path.join(tmpdir, png_files[0])
        img_standalone = Image.open(standalone_png_path)
        img_standalone.load()

        # Compare pixel size
        assert img_standalone.size == (expected_w_px, expected_h_px), (
            f"Standalone image size {img_standalone.size} != expected {(expected_w_px, expected_h_px)}"
        )

        # 7. Place in doc
        doc = PresentationDoc(
            title=f"E2E Deck - {plot_type}",
            themeId="slate",
            figures={rec["id"]: rec},
            slides=[
                SlideModel(
                    id="slide_1",
                    layout="blank",
                    objects=[
                        SlideObject(
                            id=f"obj_{plot_type}",
                            type="plot",
                            figureId=rec["id"],
                            src=final_data_url,
                            x=120,
                            y=90,
                            w=520,
                            h=390,
                            name=f"{plot_type.capitalize()} Figure",
                            alt=f"Alt text for {plot_type}",
                        )
                    ],
                )
            ],
        )

        # 8. /presentation/export PPTX
        res_pptx = client.post("/api/presentation/export", json={"format": "pptx", "doc": doc.model_dump()})
        assert res_pptx.status_code == 200, res_pptx.text
        prs = Presentation(io.BytesIO(res_pptx.content))
        assert len(prs.slides) == 1
        slide = prs.slides[0]
        # Verify picture was inserted
        pics = [s for s in slide.shapes if hasattr(s, "image")]
        assert len(pics) >= 1, "Exported PPTX does not contain the inserted figure picture"

        # 9. LibreOffice render check (skip if absent)
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        pdftoppm = shutil.which("pdftoppm")
        if not soffice or not pdftoppm:
            pytest.skip("LibreOffice or pdftoppm not installed on this system")

        pptx_path = os.path.join(tmpdir, "deck.pptx")
        with open(pptx_path, "wb") as f:
            f.write(res_pptx.content)

        # Convert PPTX to PDF via LibreOffice headless
        cmd_lo = [soffice, "--headless", "--convert-to", "pdf", pptx_path, "--outdir", tmpdir]
        res_lo = subprocess.run(cmd_lo, capture_output=True, text=True, timeout=30)
        assert res_lo.returncode == 0, f"LibreOffice conversion failed: {res_lo.stderr}"

        pdf_path = os.path.join(tmpdir, "deck.pdf")
        assert os.path.exists(pdf_path), "LibreOffice did not generate deck.pdf"

        # Convert PDF to PNG via pdftoppm
        page_prefix = os.path.join(tmpdir, "page")
        cmd_ppm = [pdftoppm, "-r", "96", "-scale-to-x", "1280", "-scale-to-y", "720", "-png", pdf_path, page_prefix]
        res_ppm = subprocess.run(cmd_ppm, capture_output=True, text=True, timeout=30)
        assert res_ppm.returncode == 0, f"pdftoppm failed: {res_ppm.stderr}"

        rendered_png = os.path.join(tmpdir, "page-1.png")
        assert os.path.exists(rendered_png), "Rendered slide PNG was not produced"

        rendered_img = Image.open(rendered_png).convert("RGB")
        assert rendered_img.size == (1280, 720)

        # Verify figure is visibly rendered in the placement area (non-white content)
        # Search area: roughly x in [100, 700], y in [80, 500]
        crop_box = (120, 90, 640, 480)
        cropped = rendered_img.crop(crop_box)
        pixels = np.asarray(cropped)
        has_non_white = bool(np.any(pixels < 235))
        assert has_non_white, f"Figure was not rendered inside placement area {crop_box} by LibreOffice"


# =====================================================================
# 3. Parity Test: Generated Script vs Endpoint Output
# =====================================================================

@pytest.mark.parametrize("plot_type", PLOT_TYPES)
def test_generated_script_vs_endpoint_parity(plot_type: str):
    """Verify generated script execution and endpoint render produce identical dimensions
    and pixel values within tolerance.
    """
    rec = make_sample_record(plot_type, w=4.0, h=3.0, dpi=150)

    # Endpoint output
    res_png = client.post("/api/figures/export", json={"record": rec, "format": "png", "dpi": 150})
    assert res_png.status_code == 200, res_png.text
    img_endpoint = Image.open(io.BytesIO(res_png.content)).convert("RGB")

    # Script output
    res_py = client.post("/api/figures/export", json={"record": rec, "format": "py", "dpi": 150})
    assert res_py.status_code == 200, res_py.text

    with tempfile.TemporaryDirectory() as tmpdir:
        py_file = os.path.join(tmpdir, "script.py")
        with open(py_file, "wb") as f:
            f.write(res_py.content)

        env = os.environ.copy()
        env["MPLBACKEND"] = "Agg"
        proc = subprocess.run([sys.executable, py_file], cwd=tmpdir, env=env, capture_output=True, text=True, timeout=20)
        assert proc.returncode == 0, proc.stderr

        out_pngs = [f for f in os.listdir(tmpdir) if f.endswith(".png")]
        assert len(out_pngs) >= 1
        img_script = Image.open(os.path.join(tmpdir, out_pngs[0])).convert("RGB")

        # 1. Identical dimensions
        assert img_script.size == img_endpoint.size == (600, 450)

        # 2. Pixel parity within rendering tolerance (mean diff < 3.0 out of 255)
        arr_ep = np.asarray(img_endpoint, dtype=np.float64)
        arr_sc = np.asarray(img_script, dtype=np.float64)
        diff = np.abs(arr_ep - arr_sc)
        mean_diff = float(np.mean(diff))
        assert mean_diff < 3.0, f"Parity diff {mean_diff:.4f} exceeds 3.0 for {plot_type}"


# =====================================================================
# 4. Composite E2E (All 6 Layouts)
# =====================================================================

@pytest.mark.parametrize("layout_id", COMPOSITE_LAYOUTS)
def test_composite_e2e_all_layouts(layout_id: str):
    """Verify all 6 composite layouts render successfully with structured panel figures."""
    from backend.figures.codegen.composite import grid

    rows, cols, cells = grid(layout_id)
    slot_count = len(cells)

    # Build composite record with panels for all slots
    panels: List[Dict[str, Any]] = []
    for slot_idx in range(slot_count):
        pt = PLOT_TYPES[slot_idx % len(PLOT_TYPES)]
        panel_rec = make_sample_record(pt, w=3.0, h=2.5, dpi=100)
        panels.append({"slot": slot_idx, "record": panel_rec})

    comp_rec = default_record("composite").model_dump()
    comp_rec["name"] = f"Composite {layout_id}"
    comp_rec["composite"].update(
        layout=layout_id,
        panels=panels,
        labels="A",
        spacing="normal",
    )
    comp_rec["export"].update(wIn=7.0, hIn=5.0, dpi=100)

    res = client.post("/api/figures/render", json={"record": comp_rec, "purpose": "preview", "dpi": 100})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is True and body["status"] == "ok"
    assert body["preview"].startswith("data:image/png;base64,")
    img = decode_png(body["preview"])
    assert img.width == 700 and img.height == 500


# =====================================================================
# 5. Python Figure E2E (All Templates)
# =====================================================================

def _get_python_templates() -> List[Dict[str, Any]]:
    return templates.list_templates()["templates"]


@pytest.mark.parametrize("tmpl", _get_python_templates(), ids=lambda t: t["id"])
def test_python_figure_e2e_templates(tmpl: Dict[str, Any]):
    """Verify all starter Python figure templates run and render cleanly."""
    res = client.post("/api/figures/python/run", json={
        "source": tmpl["source"],
        "dpi": 100,
        "purpose": "preview",
    })
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is True, f"Python template '{tmpl['id']}' failed: {body.get('stderr')}"
    assert body["status"] == "ok"
    assert body["preview"].startswith("data:image/png;base64,")
    assert body["figuresFound"] >= 1

    img = decode_png(body["preview"])
    assert img.width > 0 and img.height > 0


# =====================================================================
# 6. LaTeX Figure E2E (Skip if no TeX)
# =====================================================================

def _tex_available() -> bool:
    rt = get_runtime()
    tex_engines = rt.get("tex", {}).get("engines", {})
    return any(tex_engines.values()) or bool(shutil.which("pdflatex"))


def _get_latex_templates() -> List[Dict[str, Any]]:
    return templates.list_latex_templates()["templates"]


@pytest.mark.slow
@pytest.mark.parametrize("tmpl", _get_latex_templates(), ids=lambda t: t["id"])
def test_latex_figure_e2e_templates(tmpl: Dict[str, Any]):
    """Verify all LaTeX figure templates compile and render (skipped if no TeX compiler)."""
    if not _tex_available():
        pytest.skip("No TeX compiler installed on this system")

    res = client.post("/api/figures/latex/compile", json={
        "source": tmpl["source"],
        "mode": tmpl["mode"],
        "engine": tmpl["engine"],
        "dpi": 100,
    })
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is True, f"LaTeX compile failed for {tmpl['id']}: {body.get('message')} - {body.get('errors')}"
    assert body["status"] == "ok"
    assert body.get("pdfBase64")
    assert body.get("preview") and body["preview"].startswith("data:image/png;base64,")


# =====================================================================
# 7. Failure Cases (ok:false, human messages, intact records)
# =====================================================================

def test_failure_malformed_csv():
    """Malformed CSV input produces ok:false with human readable errors and leaves state intact."""
    # 1. Binary file upload
    res1 = client.post("/api/figures/dataset/parse", files={"file": ("binary.csv", b"\x00\x01\x02\x03not_csv\x00", "text/csv")})
    assert res1.status_code == 200
    body1 = res1.json()
    assert body1["ok"] is False and body1["status"] == "error"
    assert any(err["code"] in ("binary_file", "malformed_csv") for err in body1["errors"])
    assert any("text file" in err["message"].lower() or "csv" in err["message"].lower() for err in body1["errors"])

    # 2. Empty CSV
    res2 = client.post("/api/figures/dataset/parse", json={"text": ""})
    assert res2.status_code == 200
    body2 = res2.json()
    assert body2["ok"] is False and body2["status"] == "error"
    assert body2["errors"][0]["code"] == "empty_file"
    assert "empty" in body2["errors"][0]["message"].lower()

    # 3. Figure record carrying malformed dataset (exceeding column limits)
    rec = make_sample_record("line")
    rec["dataset"]["columns"] = [{"name": f"c{i}", "type": "number"} for i in range(300)]
    rec["dataset"]["rows"] = [[1.0] * 300]
    saved_snapshot = copy.deepcopy(rec)

    res3 = client.post("/api/figures/render", json={"record": rec, "purpose": "preview"})
    assert res3.status_code == 200
    body3 = res3.json()
    assert body3["ok"] is False and body3["status"] == "invalid"
    assert any("column" in d["message"].lower() for d in body3["diagnostics"])
    assert rec == saved_snapshot


def test_failure_log_axis_with_zeros():
    """Log-axis on non-positive values returns ok:false with human message and leaves record intact."""
    rec = make_sample_record("line")
    # Insert 0 into X data
    rec["dataset"]["rows"][0][0] = 0.0
    rec["config"]["options"]["xScale"] = "log"
    saved_snapshot = copy.deepcopy(rec)

    res = client.post("/api/figures/render", json={"record": rec, "purpose": "preview"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is False and body["status"] == "invalid"
    assert any(
        "log" in d["code"] or "<= 0" in d["message"] or "positive" in d["message"].lower()
        for d in body["diagnostics"]
    )
    # Ensure input record was not mutated
    assert rec == saved_snapshot


def test_failure_empty_data():
    """Empty dataset rows return ok:false with human message and leave record intact."""
    rec = make_sample_record("bar")
    rec["dataset"]["rows"] = []
    saved_snapshot = copy.deepcopy(rec)

    res = client.post("/api/figures/render", json={"record": rec, "purpose": "preview"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is False and body["status"] == "invalid"
    assert any("row" in d["message"].lower() or "empty" in d["message"].lower() for d in body["diagnostics"])
    assert rec == saved_snapshot


def test_failure_timeout(monkeypatch):
    """Script timing out produces ok:false with a clear human message and leaves record intact."""
    src = "import time\nimport matplotlib.pyplot as plt\nplt.plot([1, 2])\ntime.sleep(10)\n"
    rec = default_record("python").model_dump()
    rec["source"] = src
    saved_snapshot = copy.deepcopy(rec)

    # Monkeypatch runner timeout to 1 second to keep pytest fast
    real_run_python = runner.run_python

    def short_run_python(*args, **kwargs):
        kwargs["timeout"] = 1.0
        return real_run_python(*args, **kwargs)

    monkeypatch.setattr(render_core, "run_python", short_run_python)

    res = client.post("/api/figures/render", json={"record": rec, "purpose": "preview"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is False and body["status"] == "timeout"
    assert any("time" in d["message"].lower() or "stopped" in d["message"].lower() for d in body["diagnostics"])
    assert rec == saved_snapshot


def test_failure_missing_module():
    """Missing Python module produces ok:false (missing_dependency) with clear message."""
    src = "import matplotlib.pyplot as plt\nimport definitely_not_installed_package_xyz_999\nplt.plot([1, 2])\n"
    rec = default_record("python").model_dump()
    rec["source"] = src
    saved_snapshot = copy.deepcopy(rec)

    res = client.post("/api/figures/render", json={"record": rec, "purpose": "preview"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is False and body["status"] == "missing_dependency"
    assert any("definitely_not_installed_package_xyz_999" in d["message"] for d in body["diagnostics"])
    assert rec == saved_snapshot


def test_failure_missing_tex(monkeypatch):
    """Missing TeX engine produces ok:false with clear human message."""
    src = r"\documentclass{article}\begin{document}Hello\end{document}"
    # Simulate missing pdflatex compiler
    monkeypatch.setattr(latex_core, "find_engine", lambda e: None)

    res = client.post("/api/figures/latex/compile", json={"source": src, "engine": "pdflatex"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is False and body["status"] == "no_compiler"
    assert "not found on this computer" in body["message"].lower()
