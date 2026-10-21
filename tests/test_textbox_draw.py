import socket
import threading
import time
import io
import pytest
from playwright.sync_api import sync_playwright
from pptx import Presentation
from backend.exporter import EMU_PER_PIXEL, export_presentation
from backend.models import PresentationDoc

def _is_port_in_use(port: int = 8000, host: str = "127.0.0.1") -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex((host, port)) == 0

@pytest.fixture(scope="module", autouse=True)
def _ensure_server():
    server = None
    if not _is_port_in_use(8000):
        import uvicorn
        from backend.server import app
        config = uvicorn.Config(app, host="127.0.0.1", port=8000, log_level="warning")
        srv = uvicorn.Server(config)
        th = threading.Thread(target=srv.run, daemon=True)
        th.start()
        for _ in range(50):
            if _is_port_in_use(8000):
                break
            time.sleep(0.1)
        server = srv
    yield
    if server is not None:
        server.should_exit = True

def test_textbox_draw_and_pptx():
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path='/usr/bin/chromium-browser', headless=True)
        page = browser.new_page(viewport={"width": 1400, "height": 900})
        page.goto('http://127.0.0.1:8000')
        page.wait_for_load_state("networkidle")

        # Helper to clear slide objects
        def clear_slide():
            page.evaluate("""() => {
                const S = window.RPE.store;
                S.mutate('Clear objects', () => {
                    S.activeSlide().objects = [];
                });
                S.select([]);
            }""")
            time.sleep(0.1)

        def get_frame_box():
            return page.locator('#slideFrame').bounding_box()

        def to_screen_coords(frame_box, slide_x, slide_y):
            scale = page.evaluate("() => window.RPE.canvas.getScale()")
            return (frame_box['x'] + slide_x * scale, frame_box['y'] + slide_y * scale)

        # -------------------------------------------------------------------
        # Test 1: Enter draw mode, cursor is crosshair, Esc cancels before drag
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        is_drawing = page.evaluate("() => document.getElementById('slideFrame').classList.contains('drawing')")
        assert is_drawing, "Frame should have 'drawing' class after insert.textbox"

        # Press Esc to cancel
        page.keyboard.press("Escape")
        is_drawing = page.evaluate("() => document.getElementById('slideFrame').classList.contains('drawing')")
        assert not is_drawing, "Frame should not have 'drawing' class after Escape"

        # -------------------------------------------------------------------
        # Test 2: Drag in 4 directions
        # -------------------------------------------------------------------
        directions = [
            ("down-right", 100, 100, 300, 250, 100, 100, 200, 150),
            ("up-left", 400, 300, 150, 120, 150, 120, 250, 180),
            ("up-right", 100, 300, 350, 150, 100, 150, 250, 150),
            ("down-left", 400, 100, 200, 280, 200, 100, 200, 180),
        ]

        for name, sx, sy, ex, ey, exp_x, exp_y, exp_w, exp_h in directions:
            clear_slide()
            page.evaluate("() => window.RPE.commands.run('insert.textbox')")
            fb = get_frame_box()
            p1 = to_screen_coords(fb, sx, sy)
            p2 = to_screen_coords(fb, ex, ey)

            page.mouse.move(p1[0], p1[1])
            page.mouse.down()
            # Move halfway and check marquee exists, but NO ghost exists
            mid_x = (p1[0] + p2[0]) / 2
            mid_y = (p1[1] + p2[1]) / 2
            page.mouse.move(mid_x, mid_y)
            time.sleep(0.05)
            has_marquee = page.evaluate("() => !!document.querySelector('.marquee--draw')")
            has_ghost = page.evaluate("() => !!document.querySelector('.obj--ghost')")
            assert has_marquee, f"Dashed preview box should exist during {name} drag"
            assert not has_ghost, f"No shape ghost should exist during text box {name} drag"

            page.mouse.move(p2[0], p2[1])
            page.mouse.up()
            time.sleep(0.1)

            objs = page.evaluate("() => window.RPE.store.activeSlide().objects")
            assert len(objs) == 1, f"Expected 1 object after {name} drag, got {len(objs)}"
            tb = objs[0]
            assert tb["type"] == "text"
            assert tb["role"] == "text"
            assert tb["fontSize"] == 24
            assert tb["align"] == "left"
            assert tb["vAlign"] == "top"
            assert tb["stroke"] is None
            assert abs(tb["x"] - exp_x) <= 2, f"{name}: expected x={exp_x}, got {tb['x']}"
            assert abs(tb["y"] - exp_y) <= 2, f"{name}: expected y={exp_y}, got {tb['y']}"
            assert abs(tb["w"] - exp_w) <= 2, f"{name}: expected w={exp_w}, got {tb['w']}"
            assert abs(tb["h"] - exp_h) <= 2, f"{name}: expected h={exp_h}, got {tb['h']}"

            # Verify editing is active
            editing_id = page.evaluate("() => window.RPE.store.ui.editingId")
            assert editing_id == tb["id"], f"{name}: object should be in editing mode"

        # -------------------------------------------------------------------
        # Test 3: Typing works immediately
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 200, 200)
        p2 = to_screen_coords(fb, 500, 350)
        page.mouse.move(p1[0], p1[1])
        page.mouse.down()
        page.mouse.move(p2[0], p2[1])
        page.mouse.up()
        time.sleep(0.1)

        # Type immediately without clicking
        page.keyboard.type("Testing Immediate Typing 123")
        time.sleep(0.1)

        # Stop edit via Escape
        page.keyboard.press("Escape")
        time.sleep(0.1)

        saved_text = page.evaluate("() => window.RPE.store.activeSlide().objects[0].html")
        assert "Testing Immediate Typing 123" in saved_text, f"Typed text not saved, got: {saved_text}"

        # -------------------------------------------------------------------
        # Test 4: One undo step removes the text box
        # -------------------------------------------------------------------
        # Notice typing generated mutate coalesce 'edit:<id>', and before that 'Insert text box'.
        # Let's test a fresh text box with no typing to verify one undo step removes it.
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 200, 200)
        p2 = to_screen_coords(fb, 500, 350)
        page.mouse.move(p1[0], p1[1])
        page.mouse.down()
        page.mouse.move(p2[0], p2[1])
        page.mouse.up()
        time.sleep(0.1)

        assert len(page.evaluate("() => window.RPE.store.activeSlide().objects")) == 1
        page.evaluate("() => window.RPE.store.undo()")
        assert len(page.evaluate("() => window.RPE.store.activeSlide().objects")) == 0, "One undo step should remove the text box"

        # -------------------------------------------------------------------
        # Test 5: Click without drag creates default 420x64 box at click point
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 250, 180)
        page.mouse.click(p1[0], p1[1])
        time.sleep(0.1)

        objs = page.evaluate("() => window.RPE.store.activeSlide().objects")
        assert len(objs) == 1, f"Expected 1 object on click, got {len(objs)}"
        tb = objs[0]
        assert abs(tb["x"] - 250) <= 2, f"Click top-left x mismatch: expected 250, got {tb['x']}"
        assert abs(tb["y"] - 180) <= 2, f"Click top-left y mismatch: expected 180, got {tb['y']}"
        assert tb["w"] == 420
        assert tb["h"] == 64
        editing_id = page.evaluate("() => window.RPE.store.ui.editingId")
        assert editing_id == tb["id"], "Click without drag should enter edit mode"

        # -------------------------------------------------------------------
        # Test 6: Click near slide border is clamped
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 1200, 700)
        page.mouse.click(p1[0], p1[1])
        time.sleep(0.1)

        objs = page.evaluate("() => window.RPE.store.activeSlide().objects")
        assert len(objs) == 1
        tb = objs[0]
        # Slide is 1280x720. 1280 - 420 = 860, 720 - 64 = 656.
        assert tb["x"] == 860
        assert tb["y"] == 656
        assert tb["w"] == 420
        assert tb["h"] == 64

        # -------------------------------------------------------------------
        # Test 7: Sensible minimum height when dragging a shallow box
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 100, 100)
        p2 = to_screen_coords(fb, 300, 110)  # w=200, h=10
        page.mouse.move(p1[0], p1[1])
        page.mouse.down()
        page.mouse.move(p2[0], p2[1])
        page.mouse.up()
        time.sleep(0.1)

        objs = page.evaluate("() => window.RPE.store.activeSlide().objects")
        assert len(objs) == 1
        tb = objs[0]
        assert abs(tb["w"] - 200) <= 2
        assert tb["h"] >= 36, f"Height should have sensible minimum, got {tb['h']}"

        # -------------------------------------------------------------------
        # Test 8: Esc cancels during a drag
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 100, 100)
        p2 = to_screen_coords(fb, 300, 300)
        page.mouse.move(p1[0], p1[1])
        page.mouse.down()
        page.mouse.move(p2[0], p2[1])
        time.sleep(0.05)
        # Press Escape while dragging
        page.keyboard.press("Escape")
        time.sleep(0.05)
        page.mouse.up()
        time.sleep(0.1)

        objs = page.evaluate("() => window.RPE.store.activeSlide().objects")
        assert len(objs) == 0, "No object should be created if Esc was pressed during drag"
        has_marquee = page.evaluate("() => !!document.querySelector('.marquee--draw')")
        assert not has_marquee, "Marquee should be removed on Esc cancel"

        # -------------------------------------------------------------------
        # Test 9: Shape drawing still works
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('shape.insert', { shape: 'rounded' })")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 150, 150)
        p2 = to_screen_coords(fb, 350, 300)
        page.mouse.move(p1[0], p1[1])
        page.mouse.down()
        page.mouse.move((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2)
        time.sleep(0.05)
        has_ghost = page.evaluate("() => !!document.querySelector('.obj--ghost')")
        assert has_ghost, "Shape ghost should appear when drawing shape"
        page.mouse.move(p2[0], p2[1])
        page.mouse.up()
        time.sleep(0.1)

        objs = page.evaluate("() => window.RPE.store.activeSlide().objects")
        assert len(objs) == 1
        assert objs[0]["type"] == "shape"
        assert objs[0]["shape"] == "rounded"
        assert abs(objs[0]["w"] - 200) <= 2
        assert abs(objs[0]["h"] - 150) <= 2

        # -------------------------------------------------------------------
        # Test 10: Exported PPTX box geometry equals editor
        # -------------------------------------------------------------------
        clear_slide()
        page.evaluate("() => window.RPE.commands.run('insert.textbox')")
        fb = get_frame_box()
        p1 = to_screen_coords(fb, 120, 80)
        p2 = to_screen_coords(fb, 480, 260)
        page.mouse.move(p1[0], p1[1])
        page.mouse.down()
        page.mouse.move(p2[0], p2[1])
        page.mouse.up()
        time.sleep(0.1)

        ed_tb = page.evaluate("() => window.RPE.store.activeSlide().objects[0]")
        from backend.models import SlideModel, SlideObject
        doc_model = PresentationDoc(
            title="Textbox Test",
            slides=[
                SlideModel(
                    id="slide1",
                    layout="blank",
                    objects=[SlideObject.model_validate(ed_tb)]
                )
            ]
        )
        buf = export_presentation(doc_model)
        prs = Presentation(buf)
        slide = prs.slides[0]
        # Textbox shape
        pptx_tb = slide.shapes[-1]
        # Convert EMUs to logical px
        pptx_x = round(pptx_tb.left / EMU_PER_PIXEL)
        pptx_y = round(pptx_tb.top / EMU_PER_PIXEL)
        pptx_w = round(pptx_tb.width / EMU_PER_PIXEL)
        pptx_h = round(pptx_tb.height / EMU_PER_PIXEL)

        assert pptx_x == ed_tb["x"], f"PPTX left {pptx_x} != editor {ed_tb['x']}"
        assert pptx_y == ed_tb["y"], f"PPTX top {pptx_y} != editor {ed_tb['y']}"
        assert pptx_w == ed_tb["w"], f"PPTX width {pptx_w} != editor {ed_tb['w']}"
        assert pptx_h == ed_tb["h"], f"PPTX height {pptx_h} != editor {ed_tb['h']}"

        browser.close()
