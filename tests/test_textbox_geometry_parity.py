"""tests/test_textbox_geometry_parity.py

Regression test suite for textbox geometry pipeline and export parity:
1. Preview textbox dimensions faithfully converted to PPTX EMUs.
2. Prevention of automatic font reduction (normAutofit) and height collapse (spAutoFit).
3. Exact internal margins (4.5pt V, 7.5pt H) matching CSS border-box padding (6px V, 10px H).
4. Indent level conversion matching editor preview padding (48px per indent level).
5. Stroke width border-box margin adjustment.
6. Body role textboxes not forced into bullets/margins, preserving full usable width.
7. Consistent text wrapping between editor preview and exported PPTX output.
"""

import io
import json
import os
import shutil
import subprocess
import tempfile
import pytest
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

from backend.exporter import EMU_PER_PIXEL, export_presentation, to_emu
from backend.models import PresentationDoc, SlideModel, SlideObject

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FONTS_DIR = os.path.join(BASE_DIR, "fonts")


def test_textbox_dimensions_and_textbody_properties():
    """Verify exported textbox has exact EMU dimensions, wrap=square, and noAutofit."""
    obj = SlideObject(
        id="tb_geom",
        type="text",
        role="text",
        x=150.0,
        y=120.0,
        w=460.0,
        h=220.0,
        fontSize=20.0,
        vAlign="middle",
        html="<p>Test textbox geometry and bounds</p>",
    )
    doc = PresentationDoc(
        title="Geom Test",
        slides=[SlideModel(id="s1", layout="blank", objects=[obj])],
    )
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    slide = prs.slides[0]
    tb = next(sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name)

    # 1. Exact coordinate and dimension conversion
    assert tb.left == to_emu(150.0)
    assert tb.top == to_emu(120.0)
    assert tb.width == to_emu(460.0)
    assert tb.height == to_emu(220.0)

    # 2. Text body properties
    tf = tb.text_frame
    assert tf.word_wrap is True
    assert tf.auto_size == MSO_AUTO_SIZE.NONE

    body_pr = tb._element.xpath(".//a:bodyPr")[0]
    assert body_pr.get("wrap") == "square"
    # Must have <a:noAutofit/> and NOT <a:spAutoFit/> or <a:normAutofit/>
    assert body_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noAutofit") is not None
    assert body_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spAutoFit") is None
    assert body_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit") is None

    # 3. Margins: 6px V (Pt(4.5)), 10px H (Pt(7.5))
    assert tf.margin_top == Pt(4.5)
    assert tf.margin_bottom == Pt(4.5)
    assert tf.margin_left == Pt(7.5)
    assert tf.margin_right == Pt(7.5)

    # 4. Vertical alignment preserved inside the bounds
    assert tf.vertical_anchor == MSO_ANCHOR.MIDDLE


def test_body_role_preserves_usable_width_without_bullets():
    """Verify role='body' with list=None is exported as plain text without bullets or marL."""
    body_obj = SlideObject(
        id="tb_body",
        type="text",
        role="body",
        x=100.0,
        y=100.0,
        w=500.0,
        h=200.0,
        fontSize=18.0,
        list=None,
        html="<p>Regular body paragraph that should NOT have bullets or indentation.</p>",
    )
    bullet_obj = SlideObject(
        id="tb_bullet",
        type="text",
        role="body",
        x=100.0,
        y=350.0,
        w=500.0,
        h=200.0,
        fontSize=18.0,
        list="bullet",
        html="<p>Actual bullet item.</p>",
    )
    doc = PresentationDoc(
        title="Body vs Bullet",
        slides=[SlideModel(id="s1", layout="blank", objects=[body_obj, bullet_obj])],
    )
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    slide = prs.slides[0]

    # Find the two textboxes
    tbs = [sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name]
    tb_body = next(sh for sh in tbs if "Regular body" in sh.text_frame.text)
    tb_bullet = next(sh for sh in tbs if "Actual bullet" in sh.text_frame.text)

    # Body textbox should have buNone, no marL, no buChar
    body_p = tb_body.text_frame.paragraphs[0]
    body_p_pr = body_p._p.pPr
    assert body_p_pr.get("marL") is None, "Body text must not have marL"
    assert body_p_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is None
    assert body_p_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone") is not None

    # Bullet textbox should have marL, indent, and buChar
    bullet_p = tb_bullet.text_frame.paragraphs[0]
    bullet_p_pr = bullet_p._p.pPr
    assert bullet_p_pr.get("marL") is not None
    assert bullet_p_pr.get("indent") is not None
    assert bullet_p_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None


def test_indent_increases_margin_left():
    """Verify obj.indent increases margin_left by 48px per level."""
    obj0 = SlideObject(id="t0", type="text", x=50, y=50, w=400, h=100, indent=0, html="<p>Level 0</p>")
    obj1 = SlideObject(id="t1", type="text", x=50, y=160, w=400, h=100, indent=1, html="<p>Level 1</p>")
    obj2 = SlideObject(id="t2", type="text", x=50, y=270, w=400, h=100, indent=2, html="<p>Level 2</p>")

    doc = PresentationDoc(title="Indent", slides=[SlideModel(id="s1", layout="blank", objects=[obj0, obj1, obj2])])
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    slide = prs.slides[0]

    tbs = [sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name]
    t0 = next(sh for sh in tbs if "Level 0" in sh.text_frame.text)
    t1 = next(sh for sh in tbs if "Level 1" in sh.text_frame.text)
    t2 = next(sh for sh in tbs if "Level 2" in sh.text_frame.text)

    # Base margin_left is Pt(7.5) = 95250 EMUs (10px)
    assert t0.text_frame.margin_left == to_emu(10.0)
    # Level 1: 10px + 48px = 58px
    assert t1.text_frame.margin_left == to_emu(10.0 + 48.0)
    # Level 2: 10px + 96px = 106px
    assert t2.text_frame.margin_left == to_emu(10.0 + 96.0)


def test_stroke_width_border_box_margin_adjustment():
    """Verify textbox strokeWidth is accounted for in margins to maintain border-box usable width."""
    obj_plain = SlideObject(id="tp", type="text", x=50, y=50, w=400, h=100, html="<p>Plain</p>")
    obj_stroked = SlideObject(id="ts", type="text", x=50, y=160, w=400, h=100, stroke="#000000", strokeWidth=3.0, html="<p>Stroked</p>")

    doc = PresentationDoc(title="Stroke", slides=[SlideModel(id="s1", layout="blank", objects=[obj_plain, obj_stroked])])
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    slide = prs.slides[0]

    tbs = [sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name]
    tp = next(sh for sh in tbs if "Plain" in sh.text_frame.text)
    ts = next(sh for sh in tbs if "Stroked" in sh.text_frame.text)

    assert tp.text_frame.margin_left == to_emu(10.0)
    assert tp.text_frame.margin_right == to_emu(10.0)
    assert tp.text_frame.margin_top == to_emu(6.0)
    assert tp.text_frame.margin_bottom == to_emu(6.0)

    # Stroked: 10 + 3 = 13px H, 6 + 3 = 9px V
    assert ts.text_frame.margin_left == to_emu(13.0)
    assert ts.text_frame.margin_right == to_emu(13.0)
    assert ts.text_frame.margin_top == to_emu(9.0)
    assert ts.text_frame.margin_bottom == to_emu(9.0)


def test_representative_sample_deck_wrapping_parity():
    """Verify representative sample deck s1_body wraps identically without artificial overflow."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice not available on PATH")

    pdftotext = shutil.which("pdftotext")
    if not pdftotext:
        pytest.skip("pdftotext (poppler) not installed — skipping geometry parity check")

    rpe_path = os.path.join(BASE_DIR, "representative_sample_deck.rpe.json")
    with open(rpe_path, "r", encoding="utf-8") as f:
        rpe_data = json.load(f)

    doc = PresentationDoc.model_validate(rpe_data["doc"])

    with tempfile.TemporaryDirectory() as tmpdir:
        pptx_path = os.path.join(tmpdir, "slide.pptx")
        pdf_path = os.path.join(tmpdir, "slide.pdf")
        buf = export_presentation(doc)
        with open(pptx_path, "wb") as f:
            f.write(buf.getvalue())

        user_dir = os.path.join(tmpdir, "lo_profile")
        env = os.environ.copy()
        user_fonts = os.path.join(user_dir, "user", "fonts")
        os.makedirs(user_fonts, exist_ok=True)
        if os.path.isdir(FONTS_DIR):
            for root, _, files in os.walk(FONTS_DIR):
                for f in files:
                    if f.endswith((".ttf", ".otf")):
                        src = os.path.join(root, f)
                        dst = os.path.join(user_fonts, f)
                        if not os.path.exists(dst):
                            os.symlink(src, dst)
            conf_path = os.path.join(tmpdir, "fonts.conf")
            with open(conf_path, "w", encoding="utf-8") as f:
                f.write(
                    '<?xml version="1.0"?>\n'
                    '<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">\n'
                    '<fontconfig>\n'
                    '  <include ignore_missing="yes">/etc/fonts/fonts.conf</include>\n'
                    f'  <dir>{FONTS_DIR}</dir>\n'
                    f'  <dir>{user_fonts}</dir>\n'
                    '</fontconfig>\n'
                )
            env["FONTCONFIG_FILE"] = conf_path
            env["SAL_FONTPATH"] = user_fonts

        cmd = [soffice, f"-env:UserInstallation=file://{user_dir}", "--headless", "--convert-to", "pdf", pptx_path, "--outdir", tmpdir]
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)

        # Use pdftotext -bbox to get per-word bounding boxes (outputs HTML with <word> elements)
        import re as _re
        bbox_result = subprocess.run(
            [pdftotext, "-bbox", pdf_path, "-"],
            capture_output=True, text=True, timeout=30
        )
        # Each word: <word xMin="X" yMin="Y" xMax="X2" yMax="Y2">TEXT</word>
        word_pattern = _re.compile(r'<word xMin="([\d.]+)" yMin="([\d.]+)"[^>]*>([^<]+)</word>')
        raw_words = [(float(m.group(1)), float(m.group(2)), m.group(3))
                     for m in word_pattern.finditer(bbox_result.stdout)]
        # Filter to the s1_body region (same approximate bounds as before)
        body_words = [(x, y, t) for x, y, t in raw_words if 45 <= x <= 460 and 130 <= y <= 340]

        lines_dict: dict = {}
        for x, y, text in body_words:
            top_bucket = round(y, 1)
            found = None
            for k in lines_dict:
                if abs(k - top_bucket) < 3:
                    found = k
                    break
            if found is None:
                found = top_bucket
                lines_dict[found] = []
            lines_dict[found].append(text)

        sorted_tops = sorted(lines_dict.keys())
        lines = [" ".join(lines_dict[t]) for t in sorted_tops]

        assert len(lines) >= 6, f"Expected at least 6 lines in s1_body, got {len(lines)}: {lines}"
        assert not any("•" in l for l in lines), "s1_body must not contain bullet points"
        assert "This paragraph exercises multi-line text" in lines[0]
        assert "bounded" in lines[1], f"'bounded' should be on line 2, got: {lines[1]}"
        assert "emphasis" in lines[2], f"'emphasis' should be on line 3, got: {lines[2]}"


def test_user_defined_dimensions_preserved_on_short_text():
    """Verify user-defined dimensions are preserved without shrinking via <a:noAutofit/>."""
    obj = SlideObject(
        id="tb_tall",
        type="text",
        role="text",
        x=80.0,
        y=90.0,
        w=420.0,
        h=320.0,
        fontSize=18.0,
        html="<p>Short text inside tall box</p>",
    )
    doc = PresentationDoc(title="Tall Box", slides=[SlideModel(id="s1", layout="blank", objects=[obj])])
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    slide = prs.slides[0]

    tb = next(sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name)
    assert tb.width == to_emu(420.0)
    assert tb.height == to_emu(320.0)

    body_pr = tb._element.xpath(".//a:bodyPr")[0]
    assert body_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noAutofit") is not None
    assert body_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spAutoFit") is None


def test_wrapping_boundary_and_intentional_breaks():
    """Verify multiline textboxes, intentional <br> line breaks, and long words."""
    obj = SlideObject(
        id="tb_multiline",
        type="text",
        role="text",
        x=100.0,
        y=100.0,
        w=380.0,
        h=250.0,
        fontSize=18.0,
        font="Outfit",
        html="<p>First Line Intentional<br>Second Line Intentional</p><p>Third paragraph with alongwordthatdoesnotfitnormally</p>",
    )
    doc = PresentationDoc(title="Multiline Breaks", slides=[SlideModel(id="s1", layout="blank", objects=[obj])])
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    slide = prs.slides[0]

    tb = next(sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name)
    assert tb.text_frame.word_wrap is True
    assert tb.text_frame.auto_size == MSO_AUTO_SIZE.NONE

    paras = tb.text_frame.paragraphs
    assert len(paras) == 2
    # Verify intentional break in paragraph 0
    full_p0_text = "".join(r.text for r in paras[0].runs)
    assert "First Line Intentional\nSecond Line Intentional" == full_p0_text
    # Verify second paragraph has the long word
    assert "alongwordthatdoesnotfitnormally" in paras[1].text

