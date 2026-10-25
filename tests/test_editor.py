"""Comprehensive test suite for the Laurevine Rose.

Tests:
1. Canonical Document Model validation and serialization.
2. Coordinate and unit conversion math.
3. HTML-to-runs text parsing (bold, italic, underline, colors, sizes).
4. PPTX export:
   - Valid OpenXML package structure.
   - Exact physical dimensions (13.333" x 7.5").
   - Multi-slide deck generation.
   - Native text boxes with formatting and alignment.
   - Native shapes (rect, rounded, ellipse, line).
   - Embedded raster picture.
   - Native editable table.
   - Back-to-front layer order preservation.
   - Speaker notes preservation.
5. Headless LibreOffice conversion to PDF and pdftoppm image rendering.
6. FastAPI backend endpoints (/api/health, /api/presentation/export, /api/presentation/save).
"""

import io
import os
import shutil
import subprocess
import tempfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.exporter import (
    EMU_PER_PIXEL,
    export_presentation,
    parse_html_to_paragraphs,
    to_emu,
)
from backend.models import PresentationDoc, SlideModel, SlideObject
from backend.sample_deck import create_representative_deck
from backend.server import app


# 1. Canonical Document Model Validation
def test_canonical_model_validation():
    deck = create_representative_deck()
    assert deck.schemaVersion == 1
    assert len(deck.slides) == 2
    assert deck.size.w == 1280.0
    assert deck.size.h == 720.0

    # Serialization roundtrip
    data = deck.model_dump()
    assert "slides" in data
    assert len(data["slides"]) == 2

    reloaded = PresentationDoc.model_validate(data)
    assert reloaded.title == deck.title
    assert len(reloaded.slides[0].objects) == len(deck.slides[0].objects)


def test_malformed_model_rejection():
    # Missing required id in slide
    with pytest.raises(Exception):
        SlideModel.model_validate({"layout": "title"})


# 2. Coordinate Conversion
def test_coordinate_conversion():
    assert to_emu(0) == 0
    assert to_emu(1) == 9525
    # 1280 logical px at 96 DPI = 13.333333 inches = 12,192,000 EMUs
    assert to_emu(1280) == 12192000
    # 720 logical px at 96 DPI = 7.5 inches = 6,858,000 EMUs
    assert to_emu(720) == 6858000


# 3. HTML-to-Runs Parser
def test_html_to_runs_parser():
    raw_html = (
        '<p><b>Bold prefix</b> and <i>italic</i> with <span style="color: #702D3B; font-size: 24pt">colored span</span>.</p>'
        '<p><u>Second paragraph</u></p>'
    )
    paras = parse_html_to_paragraphs(raw_html, default_size=18.0)
    assert len(paras) == 2

    # First paragraph runs
    p1_runs = paras[0].runs
    assert len(p1_runs) >= 3
    assert p1_runs[0].text == "Bold prefix"
    assert p1_runs[0].bold is True

    # Check italic
    assert any(r.text == "italic" and r.italic is True for r in p1_runs)

    # Check colored span
    span_run = next(r for r in p1_runs if "colored span" in r.text)
    assert span_run.size == 24.0
    assert span_run.color is not None

    # Second paragraph underline
    p2_runs = paras[1].runs
    assert len(p2_runs) >= 1
    assert p2_runs[0].underline is True


# 4. PPTX Exporter & Representative Sample Deck
def test_representative_sample_deck_export():
    deck = create_representative_deck()
    buffer = export_presentation(deck)
    assert isinstance(buffer, io.BytesIO)
    buffer.seek(0)

    # Verify opening with python-pptx
    prs = Presentation(buffer)
    assert len(prs.slides) == 2

    # Dimensions check
    assert prs.slide_width == 12192000
    assert prs.slide_height == 6858000

    # Slide 1 verification
    s1 = prs.slides[0]
    assert s1.notes_slide.notes_text_frame.text.startswith("Presenter notes for Slide 1")

    # Check shapes on Slide 1
    s1_shape_names = [getattr(sh, "name", "") for sh in s1.shapes]
    assert len(s1.shapes) >= 7

    # Find title shape and check text
    title_shape = next(
        (sh for sh in s1.shapes if sh.has_text_frame and "Representative Research Deck" in sh.text_frame.text),
        None,
    )
    assert title_shape is not None, "Title text not found on slide 1"

    # Find picture shape (embedded PNG)
    pic_shape = next(
        (sh for sh in s1.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE),
        None,
    )
    assert pic_shape is not None, "Picture shape not found on slide 1"

    # Slide 2 verification: Shapes, Line, Table, Overlapping Layer Order
    s2 = prs.slides[1]
    assert s2.notes_slide.notes_text_frame.text.startswith("Presenter notes for Slide 2")

    # Verify table
    table_shape = next(
        (sh for sh in s2.shapes if sh.has_table),
        None,
    )
    assert table_shape is not None, "Table shape not found on slide 2"
    table = table_shape.table
    assert len(table.rows) == 5
    assert len(table.columns) == 4
    assert "Export Fidelity" in table.cell(1, 0).text

    # Verify line connector
    line_shape = next(
        (sh for sh in s2.shapes if sh.shape_type == MSO_SHAPE_TYPE.LINE),
        None,
    )
    assert line_shape is not None, "Line connector not found on slide 2"

    # Verify Layer Ordering:
    # On slide 2, objects are added in order:
    # ... s2_layer_bottom (rect), s2_layer_middle (rounded), s2_layer_top (ellipse), s2_table
    text_shapes = [sh for sh in s2.shapes if sh.has_text_frame]
    texts = [sh.text_frame.text for sh in text_shapes]

    bottom_idx = next(i for i, t in enumerate(texts) if "Layer 1: Base" in t)
    middle_idx = next(i for i, t in enumerate(texts) if "Layer 2: Middle" in t)
    top_idx = next(i for i, t in enumerate(texts) if "Layer 3: Top" in t)

    assert bottom_idx < middle_idx < top_idx, (
        f"Layer ordering violated! Expected bottom ({bottom_idx}) < middle ({middle_idx}) < top ({top_idx})"
    )


def test_all_catalogue_shapes_and_adjustable_geometry_export():
    """Verify that every shape from the catalogue exports with 100% precision,
    including adjustable shapes (arrow, arrow2, callout, callout2, chevron, cross, stars)."""
    shapes_to_test = [
        ("rect", 0.5),
        ("rounded", 0.5),
        ("ellipse", 0.5),
        ("triangle", 0.5),
        ("rtriangle", 0.5),
        ("diamond", 0.5),
        ("pentagon", 0.5),
        ("hexagon", 0.5),
        ("octagon", 0.5),
        ("parallelogram", 0.5),
        ("trapezoid", 0.5),
        ("arrow", 0.2),
        ("arrow", 0.8),
        ("arrow2", 0.3),
        ("callout", 0.1),
        ("callout", 0.9),
        ("callout2", 0.25),
        ("callout2", 0.75),
        ("chevron", 0.2),
        ("chevron", 0.8),
        ("cross", 0.1),
        ("cross", 0.9),
        ("star4", 0.2),
        ("star5", 0.5),
        ("star6", 0.8),
        ("ribbon", 0.5),
        ("brace", 0.5),
        ("scalebar", 0.5),
        ("line", 0.5),
    ]

    objects = []
    for idx, (kind, adj_val) in enumerate(shapes_to_test):
        col = idx % 5
        row = idx // 5
        obj = SlideObject(
            id=f"shp_{idx}",
            type="shape",
            shape=kind,
            x=50.0 + col * 240.0,
            y=40.0 + row * 110.0,
            w=200.0,
            h=90.0,
            fill="accent" if kind not in ("line", "brace", "scalebar") else None,
            stroke="accent",
            strokeWidth=2.0,
            adj=adj_val,
            text=f"{kind} (adj={adj_val})",
        )
        objects.append(obj)

    slide = SlideModel(
        id="shapes_slide",
        layout="blank",
        objects=objects,
    )
    doc = PresentationDoc(
        title="Shape Catalogue Test",
        slides=[slide],
    )

    buffer = export_presentation(doc)
    assert isinstance(buffer, io.BytesIO)
    buffer.seek(0)

    prs = Presentation(buffer)
    assert len(prs.slides) == 1
    out_slide = prs.slides[0]
    # Check that every shape text was exported and found on the slide
    texts = [sh.text_frame.text for sh in out_slide.shapes if sh.has_text_frame]
    for kind, adj_val in shapes_to_test:
        assert any(f"{kind} (adj={adj_val})" in t for t in texts), f"Missing shape text for {kind} (adj={adj_val})"

    # Verify that headless LibreOffice can convert these custom shapes to PDF
    soffice_path = shutil.which("libreoffice") or shutil.which("soffice")
    if soffice_path:
        with tempfile.TemporaryDirectory() as tmpdir:
            pptx_path = os.path.join(tmpdir, "shapes_test.pptx")
            with open(pptx_path, "wb") as f:
                f.write(buffer.getvalue())
            cmd = [soffice_path, "--headless", "--convert-to", "pdf", pptx_path, "--outdir", tmpdir]
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
            assert res.returncode == 0, f"LibreOffice conversion failed: {res.stderr.decode()}"
            pdf_path = os.path.join(tmpdir, "shapes_test.pdf")
            assert os.path.exists(pdf_path)
            assert os.path.getsize(pdf_path) > 1000


# 5. Headless LibreOffice Conversion & Image Rendering
def test_libreoffice_rendering():
    soffice_path = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice_path:
        pytest.skip("LibreOffice not installed in environment")

    deck = create_representative_deck()
    buffer = export_presentation(deck)

    with tempfile.TemporaryDirectory() as tmpdir:
        pptx_path = os.path.join(tmpdir, "deck.pptx")
        with open(pptx_path, "wb") as f:
            f.write(buffer.getvalue())

        # Convert to PDF via headless LibreOffice
        cmd = [soffice_path, "--headless", "--convert-to", "pdf", pptx_path, "--outdir", tmpdir]
        res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
        assert res.returncode == 0, f"LibreOffice conversion failed: {res.stderr.decode()}"

        pdf_path = os.path.join(tmpdir, "deck.pdf")
        assert os.path.exists(pdf_path), "PDF output was not created by LibreOffice"
        assert os.path.getsize(pdf_path) > 1000, "PDF output is too small"

        # If pdftoppm is available, render pages to PNG to confirm visual validity
        pdftoppm_path = shutil.which("pdftoppm")
        if pdftoppm_path:
            cmd_img = [pdftoppm_path, "-png", "-r", "100", pdf_path, os.path.join(tmpdir, "slide")]
            res_img = subprocess.run(cmd_img, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=20)
            assert res_img.returncode == 0
            pngs = [f for f in os.listdir(tmpdir) if f.endswith(".png")]
            assert len(pngs) == 2, f"Expected 2 slide images rendered, got {len(pngs)}"


# 6. FastAPI Backend Endpoints
def test_fastapi_endpoints():
    client = TestClient(app)

    # Health check
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["status"] == "ok"
    assert res.json()["connected"] is True

    # Valid export
    deck = create_representative_deck()
    payload = {"format": "pptx", "doc": deck.model_dump()}
    res_export = client.post("/api/presentation/export", json=payload)
    assert res_export.status_code == 200
    assert "attachment;" in res_export.headers.get("content-disposition", "")
    assert len(res_export.content) > 1000

    # Verify exported bytes can be read by python-pptx
    prs = Presentation(io.BytesIO(res_export.content))
    assert len(prs.slides) == 2

    # Save endpoint
    res_save = client.post("/api/presentation/save", json={"doc": deck.model_dump()})
    assert res_save.status_code == 200
    assert res_save.json()["status"] == "ok"

    # PDF export endpoint
    res_pdf = client.post("/api/presentation/export", json={"format": "pdf", "doc": deck.model_dump()})
    assert res_pdf.status_code == 200
    assert "application/pdf" in res_pdf.headers.get("content-type", "")
    assert res_pdf.content.startswith(b"%PDF-")
    assert len(res_pdf.content) > 1000

    # Unsupported format rejection
    res_bad = client.post("/api/presentation/export", json={"format": "unsupported", "doc": deck.model_dump()})
    assert res_bad.status_code == 400


# 7. Partial Selection Formatting (Font, Color, Highlight, Strike) & Justify / Itemize
def test_partial_selection_formatting_and_lists():
    html_sample = (
        '<p style="text-align: justify">Leading text with '
        '<span style="font-family: Georgia; font-size: 26pt; color: rgb(220, 20, 60); background-color: rgb(255, 255, 0)">'
        'custom partial selection'
        '</span> and <mark>highlighted term</mark> plus <s>deleted text</s>.</p>'
        '<ul>'
        '<li>First bullet point with <span style="color: blue">blue text</span></li>'
        '<li>Second bullet point</li>'
        '</ul>'
        '<ol>'
        '<li>Numbered item 1</li>'
        '<li>Numbered item 2</li>'
        '</ol>'
    )
    paras = parse_html_to_paragraphs(html_sample, default_size=18.0)
    assert len(paras) == 5

    # Paragraph 1: Justified with partial styled run
    p1 = paras[0]
    assert p1.alignment == "justify"
    partial_run = next(r for r in p1.runs if "custom partial selection" in r.text)
    assert partial_run.font == "Georgia"
    assert partial_run.size == 26.0
    assert partial_run.color == (220, 20, 60)
    assert partial_run.highlight == "FFFF00"

    mark_run = next(r for r in p1.runs if "highlighted term" in r.text)
    assert mark_run.highlight == "FFFF00"

    strike_run = next(r for r in p1.runs if "deleted text" in r.text)
    assert strike_run.strike is True

    # Paragraph 2 & 3: Bullet points
    assert paras[1].list_type == "bullet"
    assert paras[1].is_bullet is True
    assert paras[2].list_type == "bullet"

    # Paragraph 4 & 5: Numbered items
    assert paras[3].list_type == "number"
    assert paras[4].list_type == "number"


# 8. All Frontend Catalogue Shapes Export
def test_all_catalogue_shapes_export():
    from backend.models import SlideSize

    shape_kinds = [
        "rect", "rounded", "ellipse", "triangle",
        "diamond", "arrow", "line", "callout", "brace", "scalebar"
    ]
    objects = []
    for i, kind in enumerate(shape_kinds):
        objects.append(
            SlideObject(
                id=f"shp_{kind}",
                type="shape",
                shape=kind,
                x=50 + (i % 5) * 230,
                y=50 + (i // 5) * 250,
                w=200,
                h=120,
                fill="#2F6F8F",
                stroke="#702D3B",
                strokeWidth=2,
                html=f"<p>Shape: <b>{kind}</b></p>",
                fontSize=14,
                align="center",
            )
        )

    doc = PresentationDoc(
        title="All Shapes Deck",
        size=SlideSize(w=1280.0, h=720.0),
        slides=[
            SlideModel(id="slide_shapes", layout="blank", objects=objects)
        ]
    )

    buffer = export_presentation(doc)
    assert isinstance(buffer, io.BytesIO)
    prs = Presentation(buffer)
    assert len(prs.slides) == 1
    # Check that shapes were generated
    shapes = prs.slides[0].shapes
    assert len(shapes) >= 10


# 9. PDF Export End-to-End with Headless LibreOffice
def test_pdf_export_end_to_end_fidelity():
    soffice_path = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice_path:
        pytest.skip("LibreOffice not installed in environment")

    from backend.models import SlideSize

    # Create a doc with justified text, highlights, colors, itemize, and shapes
    doc = PresentationDoc(
        title="Fidelity Deck",
        size=SlideSize(w=1280.0, h=720.0),
        slides=[
            SlideModel(
                id="slide_fidelity",
                layout="blank",
                objects=[
                    SlideObject(
                        id="txt_justified",
                        type="text",
                        x=60,
                        y=60,
                        w=500,
                        h=200,
                        align="justify",
                        html=(
                            '<p style="text-align: justify">This paragraph is specifically justified across multiple lines '
                            'to verify that paragraph justification is fully preserved in the final exported PDF document. '
                            'It also contains <span style="background-color: yellow; color: red">highlighted red text</span> '
                            'via partial selection.</p>'
                        ),
                        fontSize=16,
                    ),
                    SlideObject(
                        id="txt_bullets",
                        type="text",
                        x=60,
                        y=280,
                        w=500,
                        h=200,
                        list="bullet",
                        html="<p>First item in itemize list</p><p>Second item in itemize list</p>",
                        fontSize=16,
                    ),
                    SlideObject(
                        id="shp_callout",
                        type="shape",
                        shape="callout",
                        x=600,
                        y=60,
                        w=220,
                        h=140,
                        fill="#A65C68",
                        stroke="#702D3B",
                        html="<p>Callout shape</p>",
                    ),
                    SlideObject(
                        id="shp_scalebar",
                        type="shape",
                        shape="scalebar",
                        x=600,
                        y=240,
                        w=200,
                        h=60,
                        stroke="#2F6F8F",
                        strokeWidth=3,
                        html="<p>100 nm</p>",
                    ),
                    SlideObject(
                        id="shp_brace",
                        type="shape",
                        shape="brace",
                        x=850,
                        y=60,
                        w=60,
                        h=200,
                        stroke="#702D3B",
                        strokeWidth=3,
                    ),
                ],
            )
        ],
    )

    client = TestClient(app)
    res = client.post("/api/presentation/export", json={"format": "pdf", "doc": doc.model_dump()})
    assert res.status_code == 200
    assert "application/pdf" in res.headers.get("content-type", "")
    assert res.content.startswith(b"%PDF-")
    assert len(res.content) > 5000

    # Verify pdftotext extracts the text
    with tempfile.TemporaryDirectory() as td:
        pdf_path = os.path.join(td, "fidelity.pdf")
        with open(pdf_path, "wb") as f:
            f.write(res.content)
        txt = subprocess.run(["pdftotext", pdf_path, "-"], capture_output=True).stdout.decode()
        assert "justified across" in txt
        assert "multiple lines" in txt
        assert "highlighted red text" in txt
        assert "itemize list" in txt
        assert "Callout shape" in txt
        assert "100 nm" in txt


# 10. Strict DrawingML Highlight XML Sequence Compliance
def test_drawingml_highlight_xml_sequence():
    from backend.models import SlideSize

    doc = PresentationDoc(
        title="Highlight Sequence Deck",
        size=SlideSize(w=1280.0, h=720.0),
        slides=[
            SlideModel(
                id="s_hl",
                layout="blank",
                objects=[
                    SlideObject(
                        id="txt_hl",
                        type="text",
                        x=100,
                        y=100,
                        w=400,
                        h=100,
                        html='<p>Standard <mark>yellow highlight</mark> and <span style="background-color: #FFFF00; font-family: Georgia">custom yellow</span> text.</p>',
                    )
                ],
            )
        ],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    slide = prs.slides[0]
    tb = [s for s in slide.shapes if s.has_text_frame and "yellow highlight" in s.text_frame.text][0]
    p = tb.text_frame.paragraphs[0]

    # Find the highlighted runs and verify XML structure
    found_highlight = False
    for r in p.runs:
        rpr = r._r.get_or_add_rPr()
        hl = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}highlight")
        if hl is not None:
            found_highlight = True
            # Verify highlight element contains srgbClr val="FFFF00"
            clr = hl.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            assert clr is not None
            assert clr.get("val") == "FFFF00"

            # Verify that in rpr's children, highlight comes BEFORE latin, ea, cs
            child_tags = [c.tag for c in rpr]
            hl_idx = child_tags.index(hl.tag)
            for font_tag in ("latin", "ea", "cs"):
                full_tag = f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{font_tag}"
                if full_tag in child_tags:
                    font_idx = child_tags.index(full_tag)
                    assert hl_idx < font_idx, f"<a:highlight> ({hl_idx}) must precede <a:{font_tag}> ({font_idx}) according to DrawingML schema"

    assert found_highlight, "Expected at least one run with <a:highlight>"


# 11. Brace Shape Freeform Curve Export
def test_brace_shape_freeform_export():
    from backend.models import SlideSize

    doc = PresentationDoc(
        title="Brace Curve Deck",
        size=SlideSize(w=1280.0, h=720.0),
        slides=[
            SlideModel(
                id="s_brace",
                layout="blank",
                objects=[
                    SlideObject(
                        id="shp_brace",
                        type="shape",
                        shape="brace",
                        x=100,
                        y=100,
                        w=80,
                        h=240,
                        stroke="#702D3B",
                        strokeWidth=3,
                    )
                ],
            )
        ],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    slide = prs.slides[0]
    brace_shapes = [s for s in slide.shapes if "custGeom" in s._element.xml or "pathLst" in s._element.xml]
    assert len(brace_shapes) == 1
    sp = brace_shapes[0]

    # Verify that the brace is NOT exported as MSO_SHAPE.LEFT_BRACE preset
    sp_xml = sp._element.xml
    assert "prst=\"leftBrace\"" not in sp_xml, "Brace should not export as preset leftBrace"
    # Verify freeform / custom geometry is used
    assert "custGeom" in sp_xml or "pathLst" in sp_xml, "Brace should export with custom path geometry"


# 12. Font Family Resolution, CSS Font Stacks, and Cross-Platform Fallbacks
def test_font_family_resolution_and_stacks():
    from backend.exporter import normalize_font_name, ThemeResolver
    from backend.models import SlideSize

    # Test normalize_font_name handles CSS font stacks and added fonts
    assert normalize_font_name("Georgia, 'Times New Roman', serif") == "Georgia"
    assert normalize_font_name("'Times New Roman', Times, serif") == "Times New Roman"
    assert normalize_font_name("Calibri, Carlito, sans-serif") == "Calibri"
    assert normalize_font_name("Trebuchet MS, 'Lucida Grande', sans-serif") == "Trebuchet MS"
    assert normalize_font_name("Cambria, Georgia, serif") == "Cambria"
    assert normalize_font_name("Garamond, 'EB Garamond', serif") == "Garamond"
    assert normalize_font_name("Arial, 'Helvetica Neue', sans-serif") == "Arial"
    assert normalize_font_name("Verdana, Geneva, sans-serif") == "Verdana"
    assert normalize_font_name("'Courier New', Courier, monospace") == "Courier New"
    assert normalize_font_name("Outfit, sans-serif") == "Outfit"

    # Test in export: ensure selected fonts are actually applied to runs and don't default to Outfit
    test_fonts = ["Calibri", "Georgia", "Cambria", "Garamond", "Trebuchet MS"]
    objects = []
    for i, fnt in enumerate(test_fonts):
        objects.append(
            SlideObject(
                id=f"txt_{i}",
                type="text",
                x=50,
                y=50 + i * 80,
                w=400,
                h=60,
                font=fnt,
                html=f'<p style="font-family: {fnt}">Testing font {fnt}</p>',
            )
        )

    doc = PresentationDoc(
        title="Fonts Deck",
        themeId="journal",  # default theme font is Outfit
        size=SlideSize(w=1280.0, h=720.0),
        slides=[SlideModel(id="s_fonts", layout="blank", objects=objects)],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    slide = prs.slides[0]

    for fnt in test_fonts:
        tb = [s for s in slide.shapes if s.has_text_frame and f"Testing font {fnt}" in s.text_frame.text][0]
        run = tb.text_frame.paragraphs[0].runs[0]
        assert run.font.name == fnt, f"Expected run font to be {fnt}, got {run.font.name}"
        rpr = run._r.get_or_add_rPr()
        latin = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
        assert latin is not None
        assert latin.get("typeface") == fnt
        ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        assert ea is not None
        assert ea.get("typeface") == fnt


# 14. Shape Fidelity Tests (Requirements 6.1)
def test_shape_frame_invariant():
    """1. Frame invariant: exported sp.left/top/width/height == to_emu(x/y/w/h) for all shapes (lines excepted)."""
    catalogue_shapes = [
        "rect", "rounded", "ellipse", "triangle", "rtriangle", "diamond", "pentagon",
        "hexagon", "octagon", "parallelogram", "trapezoid", "arrow", "arrow2",
        "callout", "callout2", "star4", "star5", "star6", "chevron", "ribbon",
        "cross", "brace", "scalebar"
    ]
    x, y, w, h = 123.0, 234.0, 345.0, 187.0
    for kind in catalogue_shapes:
        for adj in [0.0, 0.25, 0.5, 1.0]:
            obj = SlideObject(
                id=f"shp_{kind}_{adj}",
                type="shape",
                shape=kind,
                x=x, y=y, w=w, h=h,
                adj=adj,
                stroke="accent",
                strokeWidth=2.0,
            )
            doc = PresentationDoc(
                title="Frame Test",
                slides=[SlideModel(id="s1", layout="blank", objects=[obj])],
            )
            buf = export_presentation(doc)
            prs = Presentation(buf)
            sp = prs.slides[0].shapes[-1]
            assert sp.left == to_emu(x), f"{kind} (adj={adj}) left mismatch: {sp.left} vs {to_emu(x)}"
            assert sp.top == to_emu(y), f"{kind} (adj={adj}) top mismatch: {sp.top} vs {to_emu(y)}"
            assert sp.width == to_emu(w), f"{kind} (adj={adj}) width mismatch: {sp.width} vs {to_emu(w)}"
            assert sp.height == to_emu(h), f"{kind} (adj={adj}) height mismatch: {sp.height} vs {to_emu(h)}"


def test_outline_width_conversion():
    """2. Outline width: <a:ln w> == round(strokeWidth * 9525), stroke_only max(sw, 2), sw=None -> 2."""
    from pptx.oxml.ns import qn

    # Shape with strokeWidth 4.0
    obj1 = SlideObject(id="s1", type="shape", shape="rect", x=10, y=10, w=100, h=100, stroke="accent", strokeWidth=4.0)
    # Shape with strokeWidth None -> 2.0
    obj2 = SlideObject(id="s2", type="shape", shape="triangle", x=10, y=120, w=100, h=100, stroke="accent", strokeWidth=None)
    # Stroke-only shape with strokeWidth 1.0 -> clamped to max(1.0, 2) = 2.0
    obj3 = SlideObject(id="s3", type="shape", shape="brace", x=10, y=230, w=100, h=100, stroke="accent", strokeWidth=1.0)
    # Textbox with strokeWidth 3.0
    obj4 = SlideObject(id="s4", type="text", x=10, y=340, w=100, h=100, text="Hello", stroke="accent", strokeWidth=3.0)
    # Image fallback placeholder
    obj5 = SlideObject(id="s5", type="image", src="", x=10, y=450, w=100, h=100)

    objects = [obj1, obj2, obj3, obj4, obj5]
    doc = PresentationDoc(
        title="Outline Test",
        slides=[SlideModel(id="s_ln", layout="blank", objects=objects)],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    user_shapes = list(prs.slides[0].shapes)[-len(objects):]

    def get_ln_w(shape):
        ln = shape._element.xpath(".//a:ln")
        assert ln, f"Shape {shape} missing <a:ln>"
        return int(ln[0].get("w"))

    assert get_ln_w(user_shapes[0]) == round(4.0 * 9525)
    assert get_ln_w(user_shapes[1]) == round(2.0 * 9525)
    assert get_ln_w(user_shapes[2]) == round(2.0 * 9525)
    assert get_ln_w(user_shapes[3]) == round(3.0 * 9525)
    assert get_ln_w(user_shapes[4]) == round(1.5 * 9525)


def test_no_theme_style_on_exported_shapes():
    """3. No theme style: no <p:style> element on any exported shape, connector, or theme decor."""
    from pptx.oxml.ns import qn

    objects = [
        SlideObject(id="sh1", type="shape", shape="rect", x=10, y=10, w=100, h=100, fill="accent"),
        SlideObject(id="sh2", type="shape", shape="arrow", x=120, y=10, w=100, h=100, fill="accent"),
        SlideObject(id="sh3", type="shape", shape="line", x=230, y=10, w=100, h=100, stroke="accent"),
        SlideObject(id="sh4", type="shape", shape="brace", x=340, y=10, w=100, h=100, stroke="accent"),
    ]
    doc = PresentationDoc(
        title="Style Test",
        themeId="journal",  # has decor: rule
        slides=[SlideModel(id="s_style", layout="title", objects=objects)],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    slide = prs.slides[0]

    for sp in slide.shapes:
        style_el = sp._element.find(qn("p:style"))
        assert style_el is None, f"Found unexpected <p:style> in shape {sp.name}"
        # Also assert no outer shadow
        outer_shdw = sp._element.xpath(".//a:outerShdw")
        assert len(outer_shdw) == 0, f"Found unexpected <a:outerShdw> in shape {sp.name}"


def test_custom_geometry_commands_and_beziers():
    """4. Geometry: custGeom path has w=h=1000, callout2 and brace have real cubicBezTo / quadBezTo."""
    from backend import shape_geometry

    objects = [
        SlideObject(id="c2", type="shape", shape="callout2", x=50, y=50, w=200, h=150, adj=0.3),
        SlideObject(id="br", type="shape", shape="brace", x=300, y=50, w=100, h=300, adj=0.5),
    ]
    doc = PresentationDoc(
        title="CustGeom Test",
        slides=[
            SlideModel(
                id="s_geo",
                layout="blank",
                objects=objects,
            )
        ],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    user_shapes = list(prs.slides[0].shapes)[-len(objects):]

    # Check callout2
    sp_callout = user_shapes[0]
    path_c2 = sp_callout._element.xpath(".//a:custGeom/a:pathLst/a:path")
    assert len(path_c2) == 1
    assert path_c2[0].get("w") == "1000" and path_c2[0].get("h") == "1000"
    cubics = sp_callout._element.xpath(".//a:custGeom/a:pathLst/a:path/a:cubicBezTo")
    assert len(cubics) >= 5, f"Expected cubic curves in callout2, found {len(cubics)}"

    # Check brace
    sp_brace = user_shapes[1]
    path_br = sp_brace._element.xpath(".//a:custGeom/a:pathLst/a:path")
    assert len(path_br) == 1
    assert path_br[0].get("w") == "1000" and path_br[0].get("h") == "1000"
    quads = sp_brace._element.xpath(".//a:custGeom/a:pathLst/a:path/a:quadBezTo")
    assert len(quads) >= 4, f"Expected quad curves in brace, found {len(quads)}"


def test_adjusted_geometry_survives_export_and_endpoint():
    """5. Adjusted geometry survives export: same shape at adj=0.2 vs 0.8 has different points, frame unchanged."""
    from fastapi.testclient import TestClient
    from backend.server import app

    obj_02 = SlideObject(id="a1", type="shape", shape="arrow", x=100, y=100, w=200, h=100, adj=0.2)
    obj_08 = SlideObject(id="a2", type="shape", shape="arrow", x=100, y=100, w=200, h=100, adj=0.8)

    objects = [obj_02, obj_08]
    doc = PresentationDoc(
        title="Adjust Export Test",
        slides=[SlideModel(id="s_adj", layout="blank", objects=objects)],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    user_shapes = list(prs.slides[0].shapes)[-len(objects):]
    sp1, sp2 = user_shapes[0], user_shapes[1]

    # Frames must be identical
    assert sp1.left == sp2.left and sp1.top == sp2.top
    assert sp1.width == sp2.width and sp1.height == sp2.height

    # Path XML must differ
    xml1 = sp1._element.xpath(".//a:custGeom/a:pathLst/a:path")[0]
    xml2 = sp2._element.xpath(".//a:custGeom/a:pathLst/a:path")[0]
    pts1 = [pt.attrib for pt in xml1.xpath(".//a:pt")]
    pts2 = [pt.attrib for pt in xml2.xpath(".//a:pt")]
    assert pts1 != pts2, "Paths for adj=0.2 and adj=0.8 must differ"

    # Test through HTTP endpoint
    client = TestClient(app)
    resp = client.post("/api/presentation/export", json={"format": "pptx", "doc": doc.model_dump()})
    assert resp.status_code == 200
    prs_http = Presentation(io.BytesIO(resp.content))
    assert len(prs_http.slides) == 1


def test_rounded_corner_radius_adjustment():
    """6. Rounded rectangle: sp.adjustments[0] == min(0.5, radius / min(w, h))."""
    obj1 = SlideObject(id="r1", type="shape", shape="rounded", x=50, y=50, w=200, h=100, cornerRadius=16)
    obj2 = SlideObject(id="r2", type="shape", shape="rounded", x=50, y=200, w=200, h=100, cornerRadius=80)

    objects = [obj1, obj2]
    doc = PresentationDoc(
        title="Rounded Test",
        slides=[SlideModel(id="s_rnd", layout="blank", objects=objects)],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    user_shapes = list(prs.slides[0].shapes)[-len(objects):]
    sp1 = user_shapes[0]
    sp2 = user_shapes[1]

    # 16 / min(200, 100) = 0.16
    assert abs(sp1.adjustments[0] - 0.16) < 0.001
    # 80 / 100 = 0.8 -> capped at 0.5
    assert abs(sp2.adjustments[0] - 0.5) < 0.001


def test_text_color_regression_without_p_style():
    """7. Text color regression: filled shape with text (with/without color) retains explicit run colors."""
    from pptx.dml.color import RGBColor

    obj_with_color = SlideObject(
        id="t1", type="shape", shape="rect", x=50, y=50, w=200, h=100,
        fill="accent", color="#00FF00", text="Green Text"
    )
    obj_default_color = SlideObject(
        id="t2", type="shape", shape="arrow", x=50, y=200, w=200, h=100,
        fill="accent", text="White Text"
    )

    objects = [obj_with_color, obj_default_color]
    doc = PresentationDoc(
        title="Text Color Test",
        slides=[SlideModel(id="s_tc", layout="blank", objects=objects)],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    user_shapes = list(prs.slides[0].shapes)[-len(objects):]
    sp1 = user_shapes[0]
    sp2 = user_shapes[1]

    run1 = sp1.text_frame.paragraphs[0].runs[0]
    assert run1.font.color.rgb == RGBColor(0, 255, 0)

    run2 = sp2.text_frame.paragraphs[0].runs[0]
    # For filled shape, default text color resolves to white
    assert run2.font.color.rgb == RGBColor(255, 255, 255)


def test_table_resolution_parity_fixture():
    """Table resolution parity fixture: Python resolve_cell matches table_resolve.json exactly."""
    import json
    with open("tests/fixtures/table_resolve.json", "r") as f:
        fixture = json.load(f)

    obj = SlideObject(**fixture["tableObject"])
    from backend.exporter import (
        compute_table_geometry,
        compute_table_tints,
        resolve_table_border,
        resolve_cell,
        ThemeResolver,
    )
    geom = compute_table_geometry(obj)
    assert geom == fixture["geometry"]

    for theme_id, expected in fixture["themes"].items():
        doc = PresentationDoc(themeId=theme_id)
        resolver = ThemeResolver(doc)
        is_dark = expected["dark"]
        band_tint, line_color = compute_table_tints(expected["slideBg"], is_dark, resolver)
        assert band_tint == expected["tints"]["bandTint"]
        assert line_color == expected["tints"]["lineColor"]

        border_w, border_col = resolve_table_border(obj, expected["slideBg"], is_dark, resolver)
        assert border_w == expected["border"]["width"]
        assert border_col == expected["border"]["color"]

        for cell_key, expected_cell in expected["resolvedCells"].items():
            r, c = (int(x) for x in cell_key.split(","))
            res = resolve_cell(obj, r, c, expected["slideBg"], is_dark, resolver)
            assert res == expected_cell, f"Mismatch for {theme_id} at {cell_key}: {res} != {expected_cell}"


def test_legacy_table_export_fidelity():
    """Legacy tables (no new fields) export with opaque tints, correct tblPr, margins, anchor, and borders."""
    from backend.exporter import export_presentation
    from pptx import Presentation

    cells = [["H1", "H2"], ["D1", "D2"], ["D3", "D4"]]
    legacy_obj = SlideObject(
        id="t_legacy",
        type="table",
        x=50, y=50, w=400, h=150,
        cells=cells,
        headerRow=True,
        banded=True,
    )

    for theme_id, expected_band, expected_line in [
        ("journal", "F8F5F5", "DBDBDB"),
        ("midnight", "292C33", "47494F"),
    ]:
        doc = PresentationDoc(
            themeId=theme_id,
            slides=[SlideModel(id=f"s_{theme_id}", layout="blank", objects=[legacy_obj])],
        )
        buf = export_presentation(doc)
        prs = Presentation(buf)
        tbl_shape = prs.slides[0].shapes[-1]
        tbl = tbl_shape.table._tbl

        # tblPr
        tblPr = tbl.tblPr
        assert tblPr.get("firstRow") == "1"
        assert tblPr.get("bandRow") == "0"
        style_el = tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
        assert style_el is not None
        assert style_el.text == "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

        # Header cell (0, 0): accent fill, white text, anchor ctr
        c00 = tbl_shape.table.cell(0, 0)._tc.tcPr
        assert c00.get("marL") == "95250"
        assert c00.get("marR") == "95250"
        assert c00.get("marT") == "38100"
        assert c00.get("marB") == "38100"
        assert c00.get("anchor") == "ctr"

        # Check ln* before fill in tcPr
        children = [c.tag.split("}")[-1] for c in c00]
        fill_idx = next(i for i, tag in enumerate(children) if "Fill" in tag)
        ln_indices = [i for i, tag in enumerate(children) if tag.startswith("ln")]
        assert len(ln_indices) == 4, f"Expected 4 ln borders, found {ln_indices}"
        assert all(i < fill_idx for i in ln_indices), "All ln* must strictly precede fill in tcPr"

        # Check line color in border
        lnL = c00.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lnL")
        clr = lnL.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        assert clr is not None and clr.get("val") == expected_line

        # Check banded row (r=2, c=0): bandTint
        c20 = tbl_shape.table.cell(2, 0)._tc.tcPr
        fill_20 = c20.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
        assert fill_20 is not None
        clr_20 = fill_20.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        assert clr_20 is not None and clr_20.get("val") == expected_band

        # Check unbanded row (r=1, c=0): noFill
        c10 = tbl_shape.table.cell(1, 0)._tc.tcPr
        assert c10.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noFill") is not None


def test_customized_table_export_xml_and_geometry():
    """Fully customized table: non-uniform sizes, mixed formatting, empty cells, borders, exact EMU sums."""
    from backend.exporter import export_presentation
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    cells = [["Title A", "Title B", ""], ["Data 1", "", "Data 3"], ["Data 4", "Data 5", "Data 6"]]
    col_w = [120.0, 240.0, 180.0]  # sum = 540
    row_h = [60.0, 40.0, 80.0]     # sum = 180

    cell_fmt = {
        "0,0": {"fill": "accent", "bold": True, "align": "center", "vAlign": "top", "fontSize": 22.0},
        "0,1": {"fill": None, "italic": True, "vAlign": "bottom", "font": "Georgia", "color": "#00FF00"},
        "0,2": {"fill": "#123456"},  # empty cell with custom fill
        "1,0": {"underline": True, "strike": True, "align": "right"},
        "1,1": {"fontSize": 10.0},    # empty cell with custom fontSize
        "2,1": {"fill": "accent", "bold": False, "vAlign": "middle"},
    }
    table_border = {"width": 3.0, "color": "#AABBCC"}

    obj = SlideObject(
        id="t_custom",
        type="table",
        x=100.0, y=100.0, w=540.0, h=180.0,
        cells=cells,
        colW=col_w,
        rowH=row_h,
        cellFmt=cell_fmt,
        tableBorder=table_border,
        headerRow=True,
        banded=True,
    )

    doc = PresentationDoc(
        themeId="journal",
        slides=[SlideModel(id="s_custom", layout="blank", objects=[obj])],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    tbl_shape = prs.slides[0].shapes[-1]
    tbl = tbl_shape.table

    # 1. Geometry and exact EMU sums
    grid_cols = [int(col.get("w")) for col in tbl._tbl.tblGrid.xpath("a:gridCol")]
    assert sum(grid_cols) == tbl_shape.width
    assert tbl_shape.width == round(540.0 * 9525)

    tr_heights = [int(tr.get("h")) for tr in tbl._tbl.xpath("a:tr")]
    assert sum(tr_heights) == tbl_shape.height
    assert tbl_shape.height == round(180.0 * 9525)

    # 2. Check borders and tcPr schema ordering across all cells
    for r in range(3):
        for c in range(3):
            tcPr = tbl.cell(r, c)._tc.tcPr
            children = [child.tag.split("}")[-1] for child in tcPr]
            ln_indices = [i for i, tag in enumerate(children) if tag.startswith("ln")]
            fill_indices = [i for i, tag in enumerate(children) if "Fill" in tag]
            assert len(ln_indices) == 4, f"Cell ({r},{c}) missing borders"
            assert len(fill_indices) == 1, f"Cell ({r},{c}) missing fill"
            assert max(ln_indices) < min(fill_indices), f"Cell ({r},{c}) ln* must strictly precede fill"

            # Check 3px border width and color
            lnL = tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lnL")
            assert lnL.get("w") == str(round(3.0 * 9525))
            clr = lnL.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            assert clr is not None and clr.get("val") == "AABBCC"

    # 3. Check anchors
    assert tbl.cell(0, 0)._tc.tcPr.get("anchor") == "t"
    assert tbl.cell(0, 1)._tc.tcPr.get("anchor") == "b"
    assert tbl.cell(2, 1)._tc.tcPr.get("anchor") == "ctr"

    # 4. Check text runs & empty cells
    # (0, 0)
    p00 = tbl.cell(0, 0).text_frame.paragraphs[0]
    assert p00.runs[0].font.bold is True
    assert p00.runs[0].font.size == Pt(22)
    # (0, 1)
    p01 = tbl.cell(0, 1).text_frame.paragraphs[0]
    assert p01.runs[0].font.italic is True
    assert p01.runs[0].font.name == "Georgia"
    assert p01.runs[0].font.color.rgb == RGBColor(0, 255, 0)
    # (0, 2) empty cell with endParaRPr
    p02 = tbl.cell(0, 2).text_frame.paragraphs[0]
    assert len(p02.runs) == 0
    end_rpr02 = p02._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr")
    assert end_rpr02 is not None
    # (1, 0) underline and strike
    p10 = tbl.cell(1, 0).text_frame.paragraphs[0]
    assert p10.runs[0].font.underline is True
    assert p10.runs[0]._r.rPr.get("strike") == "sngStrike"
    # (1, 1) empty cell with fontSize 10 -> sz=1000
    p11 = tbl.cell(1, 1).text_frame.paragraphs[0]
    end_rpr11 = p11._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr")
    assert end_rpr11 is not None and end_rpr11.get("sz") == "1000"


def test_table_no_header_no_banded_styling():
    """headerRow=False and banded=False produce no header fill and no banding tint."""
    from backend.exporter import export_presentation
    from pptx import Presentation

    cells = [["R0C0", "R0C1"], ["R1C0", "R1C1"], ["R2C0", "R2C1"]]
    obj = SlideObject(
        id="t_plain",
        type="table",
        x=50, y=50, w=300, h=150,
        cells=cells,
        headerRow=False,
        banded=False,
    )
    doc = PresentationDoc(
        themeId="journal",
        slides=[SlideModel(id="s_plain", layout="blank", objects=[obj])],
    )
    buf = export_presentation(doc)
    prs = Presentation(buf)
    tbl_shape = prs.slides[0].shapes[-1]

    # Row 0 and Row 2 should both have noFill
    for r in [0, 1, 2]:
        for c in [0, 1]:
            tcPr = tbl_shape.table.cell(r, c)._tc.tcPr
            assert tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noFill") is not None


# 25. Per-paragraph alignment: PPTX algn attrs reflect individual paragraph styles,
#     character formatting is untouched, and the reset-all path strips per-para overrides.
def test_per_paragraph_alignment_export():
    import re as _re
    from pptx.util import Pt

    html = (
        '<p style="text-align: justify;">Para 1 with <b>bold run</b> and normal.</p>'
        '<p style="text-align: center;">Para 2 with <span style="font-size: 24pt;">big run</span>.</p>'
        '<p style="text-align: right;">Para 3 plain.</p>'
    )
    obj = SlideObject(
        id="tb_align", type="text",
        x=100, y=100, w=500, h=300,
        html=html, align="left", fontSize=18,
    )
    doc = PresentationDoc(title="Per-Para Align", slides=[SlideModel(id="s1", layout="blank", objects=[obj])])
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))

    tb = next(s for s in prs.slides[0].shapes if s.has_text_frame and not s.is_placeholder and "Rectangle" not in s.name)
    content_paras = [p for p in tb.text_frame.paragraphs if p.text.strip()]
    assert len(content_paras) == 3, f"Expected 3 content paras, got {len(content_paras)}"
    p1, p2, p3 = content_paras

    # alignment attributes
    assert p1._p.pPr.get("algn") == "just", f"P1 algn: {p1._p.pPr.get('algn')}"
    assert p2._p.pPr.get("algn") == "ctr",  f"P2 algn: {p2._p.pPr.get('algn')}"
    assert p3._p.pPr.get("algn") == "r",    f"P3 algn: {p3._p.pPr.get('algn')}"

    # character formatting preserved
    bold_runs = [r for r in p1.runs if r.font.bold]
    assert bold_runs and "bold run" in bold_runs[0].text, "Bold run missing/wrong in P1"

    big_runs = [r for r in p2.runs if r.font.size and abs(r.font.size - Pt(24)) < 5000]
    assert big_runs, "24pt run missing from P2"

    # ------- reset-all path: strip per-para styles, box align='left' → algn=l -------
    def _strip_para_align(s):
        def _clean_style(m):
            before, style_attr = m.group(1), m.group(2)
            new = _re.sub(r'\s*text-align\s*:[^;"]+(;|(?="))', '', style_attr)
            body = _re.sub(r'style="([^"]*)"', r'\1', new).strip()
            return before + (f' style="{body}"' if body.strip().rstrip(';') else '')
        s = _re.sub(r'(<p\b[^>]*?)(\s+style="[^"]*")', _clean_style, s)
        s = _re.sub(r'(<p\b[^>]*?)\s+align="[^"]*"', r'\1', s)
        return s

    stripped = _strip_para_align(html)
    obj2 = SlideObject(id="tb2", type="text", x=100, y=100, w=500, h=300,
                       html=stripped, align="left", fontSize=18)
    doc2 = PresentationDoc(title="Reset", slides=[SlideModel(id="s2", layout="blank", objects=[obj2])])
    buf2 = export_presentation(doc2)
    prs2 = Presentation(io.BytesIO(buf2.getvalue()))
    tb2 = next(s for s in prs2.slides[0].shapes if s.has_text_frame and not s.is_placeholder and "Rectangle" not in s.name)
    for i, p in enumerate(p for p in tb2.text_frame.paragraphs if p.text.strip()):
        algn = p._p.pPr.get("algn")
        assert algn in (None, "l"), f"After reset, para {i+1} algn={algn!r} (expected l or None)"


def test_list_marker_formatting_and_alignment():
    import shutil
    import subprocess

    obj = SlideObject(
        id="txt_list_test",
        type="text",
        x=50,
        y=50,
        w=500,
        h=400,
        list="number",
        html='<p><span style="font-family: Georgia; font-size: 28pt; color: rgb(220, 20, 60);">First item Georgia 28pt</span></p><p><span style="font-family: Georgia; font-size: 14pt;">Second item 14pt</span></p>',
        fontSize=18,
    )
    doc = PresentationDoc(title="List Test", slides=[SlideModel(id="s1", layout="blank", objects=[obj])])
    buf = export_presentation(doc)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    s = prs.slides[0]
    tb = next(sh for sh in s.shapes if sh.has_text_frame and not sh.is_placeholder and "Rectangle" not in sh.name)
    p0 = tb.text_frame.paragraphs[0]
    p1 = tb.text_frame.paragraphs[1]

    # Verify marL and indent derived from run font size
    assert p0._p.pPr.get("marL") == str(int(round(28 * 1.35 * 12700)))
    assert p0._p.pPr.get("indent") == str(int(round(-28 * 1.35 * 12700)))
    assert p1._p.pPr.get("marL") == str(int(round(14 * 1.35 * 12700)))
    assert p1._p.pPr.get("indent") == str(int(round(-14 * 1.35 * 12700)))

    # Verify valid DrawingML pPr child sequence order
    tag_names = [child.tag.split("}")[-1] for child in p0._p.pPr]
    assert "buClrTx" in tag_names
    assert "buSzTx" in tag_names
    assert "buFontTx" in tag_names
    assert "buAutoNum" in tag_names
    idx_clr = tag_names.index("buClrTx")
    idx_sz = tag_names.index("buSzTx")
    idx_font = tag_names.index("buFontTx")
    idx_num = tag_names.index("buAutoNum")
    assert idx_clr < idx_sz < idx_font < idx_num

    # LibreOffice render verification: confirm the PDF was produced
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory() as tmpdir:
            pptx_path = os.path.join(tmpdir, "test.pptx")
            with open(pptx_path, "wb") as f:
                f.write(buf.getvalue())
            result = subprocess.run([soffice, "--headless", "--convert-to", "pdf", pptx_path, "--outdir", tmpdir], check=True, capture_output=True)
            pdf_path = os.path.join(tmpdir, "test.pdf")
            assert os.path.isfile(pdf_path), "LibreOffice failed to produce a PDF"


def test_format_painter_export():
    # 1. Whole box format painter: stale inline spans stripped, uniform object props applied
    # Box 2 after format painter from Box 1 (fontSize=24, font='Outfit', bold=True, color='#C00000')
    # HTML has stale spans stripped to plain paragraph
    obj_whole = SlideObject(
        id="box_target",
        type="text",
        x=50, y=50, w=400, h=100,
        fontSize=24.0,
        font="Outfit",
        bold=True,
        color="#C00000",
        align="right",
        list="bullet",
        html="<p>Uniform text</p>"
    )
    paras_whole = parse_html_to_paragraphs(
        obj_whole.html,
        default_bold=obj_whole.bold,
        default_size=obj_whole.fontSize,
        default_font=obj_whole.font,
    )
    assert len(paras_whole) == 1
    run_whole = paras_whole[0].runs[0]
    assert run_whole.text == "Uniform text"
    assert run_whole.bold is True
    assert run_whole.size == 24.0
    assert run_whole.font == "Outfit"

    # 2. Character format painter: paint bold-red-24pt word onto a word only
    # Range styling applied only to that word, surrounding runs preserved
    html_range = '<p>Normal text <span style="font-size: 24pt; font-weight: bold; color: rgb(255, 0, 0)">painted</span> untouched text</p>'
    paras_range = parse_html_to_paragraphs(html_range, default_size=18.0, default_bold=False)
    assert len(paras_range) == 1
    runs = paras_range[0].runs
    assert len(runs) >= 3

    painted_run = next(r for r in runs if "painted" in r.text)
    assert painted_run.size == 24.0
    assert painted_run.bold is True
    assert painted_run.color == (255, 0, 0)

    normal_run = next(r for r in runs if "Normal" in r.text)
    assert normal_run.size == 18.0
    assert normal_run.bold is False

    untouched_run = next(r for r in runs if "untouched" in r.text)
    assert untouched_run.size == 18.0
    assert untouched_run.bold is False


