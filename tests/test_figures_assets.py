"""Image assets of persisted documents (contract 3e): backend/assets.py, PresentationDoc, save / export round trip."""
import base64
import io

import numpy as np
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.assets import resolve_assets
from backend.figures.models import default_record
from backend.models import PresentationDoc, upgrade_document
from backend.server import app

PIX = "data:image/png;base64,iVBORw0KGgo="


def png_url(w, h, seed):
    """Noise PNG (does not compress), as a data URL."""
    arr = np.random.default_rng(seed).integers(0, 256, (h, w, 3), dtype=np.uint8)
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, "PNG", compress_level=1)
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


def obj(oid, typ="image", src=None, **extra):
    return {"id": oid, "type": typ, "x": 10, "y": 10, "w": 320, "h": 180, "src": src, **extra}


def doc_with(objects, figures=None, assets=None):
    d = {"title": "Assets", "size": {"w": 1280, "h": 720}, "schemaVersion": 2, "figures": figures or {},
         "slides": [{"id": "s1", "layout": "blank", "objects": objects}]}
    if assets is not None:
        d["assets"] = assets
    return d


# ---------- pure resolver ----------

def test_resolve_replaces_refs_and_drops_assets_key():
    doc = doc_with([obj("a", src="asset:x1"), obj("b", src="asset:x1"), obj("c", src=None)], assets={"x1": PIX})
    out = resolve_assets(doc)
    assert "assets" not in out
    assert [o["src"] for o in out["slides"][0]["objects"]] == [PIX, PIX, None]
    assert doc["slides"][0]["objects"][0]["src"] == "asset:x1" and "assets" in doc, "input is not modified"


def test_resolve_figure_preview_and_dangling_reference():
    rec = default_record("structured", "line", id="fig_1").model_dump()
    rec["render"]["preview"] = "asset:p1"
    rec2 = default_record("structured", "bar", id="fig_2").model_dump()
    rec2["render"]["preview"] = "asset:gone"
    doc = doc_with([obj("a", "plot", "asset:missing", figureId="fig_1")], {"fig_1": rec, "fig_2": rec2}, {"p1": PIX})
    out = resolve_assets(doc)
    assert out["figures"]["fig_1"]["render"]["preview"] == PIX
    assert out["figures"]["fig_2"]["render"]["preview"] is None
    assert out["slides"][0]["objects"][0]["src"] is None and out["slides"][0]["objects"][0]["figureId"] == "fig_1"


def test_resolve_leaves_plain_documents_alone():
    doc = doc_with([obj("a", src=PIX)])
    assert resolve_assets(doc) is doc


@pytest.mark.parametrize("assets", [[], {"bad id": PIX}, {"ok": "http://example.com/x.png"}, {"ok": 5}])
def test_resolve_rejects_bad_assets(assets):
    with pytest.raises(ValueError):
        resolve_assets(doc_with([obj("a", src="asset:ok")], assets=assets))


# ---------- models ----------

def test_presentation_doc_reads_assets_and_keeps_the_figure_link():
    rec = default_record("structured", "scatter", id="fig_a").model_dump()
    rec["render"].update({"status": "ok", "preview": "asset:p"})
    doc = PresentationDoc.model_validate(doc_with([obj("o1", "plot", "asset:f", figureId="fig_a")], {"fig_a": rec}, {"p": PIX, "f": PIX}))
    assert doc.slides[0].objects[0].src == PIX and doc.slides[0].objects[0].figureId == "fig_a"
    assert doc.figures["fig_a"].render.preview == PIX
    assert upgrade_document(doc).slides[0].objects[0].src == PIX


# ---------- HTTP round trip of a ~20 MB document ----------

@pytest.fixture(scope="module")
def big_doc():
    """Three 5 MB noise pictures (about 20 MB of JSON), one of them used twice, one as a figure picture with a preview."""
    urls = [png_url(1700, 1000, s) for s in (1, 2, 3)]
    assets = {f"a{i}": u for i, u in enumerate(urls)}
    rec = default_record("structured", "line", id="fig_big").model_dump()
    rec["render"].update({"status": "ok", "preview": "asset:a2"})
    objects = [obj("i0", src="asset:a0"), obj("i1", src="asset:a0"), obj("i2", src="asset:a1"), obj("p0", "plot", "asset:a2", figureId="fig_big")]
    return doc_with(objects, {"fig_big": rec}, assets), urls


def test_big_document_save_round_trip(big_doc):
    doc, urls = big_doc
    assert sum(len(u) for u in urls) > 20_000_000
    res = TestClient(app).post("/api/presentation/save", json={"doc": doc})
    assert res.status_code == 200 and res.json()["status"] == "ok"


def test_big_document_export_embeds_each_picture(big_doc):
    doc, urls = big_doc
    res = TestClient(app).post("/api/presentation/export", json={"format": "pptx", "doc": doc})
    assert res.status_code == 200
    prs = Presentation(io.BytesIO(res.content))
    pics = [s for sl in prs.slides for s in sl.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 4
    sizes = {Image.open(io.BytesIO(p.image.blob)).size for p in pics}
    assert sizes == {(1700, 1000)}, "the pictures keep their pixels"
    assert len(res.content) > 14_000_000, "the 3 distinct pictures are all in the file (the repeated one is stored once)"


def test_save_rejects_a_broken_assets_block(big_doc):
    doc, _ = big_doc
    bad = {**doc, "assets": {"a0": "javascript:alert(1)"}}
    res = TestClient(app).post("/api/presentation/save", json={"doc": bad})
    assert res.status_code in (400, 422)


def test_dangling_reference_exports_without_picture():
    doc = doc_with([obj("i0", src="asset:nope")], assets={})
    res = TestClient(app).post("/api/presentation/export", json={"format": "pptx", "doc": doc})
    assert res.status_code == 200
    prs = Presentation(io.BytesIO(res.content))
    assert not [s for sl in prs.slides for s in sl.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
