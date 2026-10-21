"""
Tests for dedicated line shape geometry, creation, serialization, and PPTX export.
"""
import io
import json
import shutil
import subprocess
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend import shape_geometry
from backend.exporter import export_presentation, ThemeResolver, to_emu
from backend.models import PresentationDoc, SlideModel, SlideObject


# ============================================================================
# 1. Geometry and Bounds Tests
# ============================================================================

def test_line_geometry_explicit_endpoints():
    # Diagonal line
    obj = SlideObject(id="l1", type="shape", shape="line", x=10, y=20, w=100, h=50, x1=10, y1=20, x2=110, y2=70)
    assert shape_geometry.is_line(obj) is True
    x1, y1, x2, y2 = shape_geometry.line_geometry(obj)
    assert (x1, y1, x2, y2) == (10.0, 20.0, 110.0, 70.0)

    bounds = shape_geometry.line_bounds(x1, y1, x2, y2)
    assert bounds == (10.0, 20.0, 100.0, 50.0)


def test_line_geometry_reversed_endpoints():
    # Reversed drawing direction: right-to-left, bottom-to-top
    obj = SlideObject(id="l2", type="shape", shape="line", x1=300, y1=200, x2=100, y2=50)
    x1, y1, x2, y2 = shape_geometry.line_geometry(obj)
    assert (x1, y1, x2, y2) == (300.0, 200.0, 100.0, 50.0)

    x, y, w, h = shape_geometry.line_bounds(x1, y1, x2, y2)
    assert x == 100.0
    assert y == 50.0
    assert w == 200.0
    assert h == 150.0


def test_line_geometry_zero_dimensions():
    # Horizontal line (height = 0)
    h_obj = SlideObject(id="lh", type="shape", shape="line", x1=50, y1=100, x2=250, y2=100)
    x1, y1, x2, y2 = shape_geometry.line_geometry(h_obj)
    assert (x1, y1, x2, y2) == (50.0, 100.0, 250.0, 100.0)
    x, y, w, h = shape_geometry.line_bounds(x1, y1, x2, y2)
    assert (x, y, w, h) == (50.0, 100.0, 200.0, 0.0)

    # Vertical line (width = 0)
    v_obj = SlideObject(id="lv", type="shape", shape="line", x1=150, y1=40, x2=150, y2=340)
    x1, y1, x2, y2 = shape_geometry.line_geometry(v_obj)
    assert (x1, y1, x2, y2) == (150.0, 40.0, 150.0, 340.0)
    x, y, w, h = shape_geometry.line_bounds(x1, y1, x2, y2)
    assert (x, y, w, h) == (150.0, 40.0, 0.0, 300.0)

    # Zero-length line (point)
    pt_obj = SlideObject(id="lp", type="shape", shape="line", x1=80, y1=80, x2=80, y2=80)
    x1, y1, x2, y2 = shape_geometry.line_geometry(pt_obj)
    assert (x1, y1, x2, y2) == (80.0, 80.0, 80.0, 80.0)
    assert shape_geometry.line_bounds(x1, y1, x2, y2) == (80.0, 80.0, 0.0, 0.0)


def test_line_geometry_legacy_fallback():
    # Legacy object without x1/y1/x2/y2 (from earlier document format)
    legacy = SlideObject(id="old_line", type="shape", shape="line", x=100, y=200, w=400, h=24)
    x1, y1, x2, y2 = shape_geometry.line_geometry(legacy)
    assert x1 == 100.0
    assert y1 == 212.0  # y + h/2
    assert x2 == 500.0  # x + w
    assert y2 == 212.0


def test_is_line_detection():
    assert shape_geometry.is_line(SlideObject(id="1", type="line")) is True
    assert shape_geometry.is_line(SlideObject(id="2", type="shape", shape="line")) is True
    assert shape_geometry.is_line(SlideObject(id="3", type="shape", shape="rect")) is False
    assert shape_geometry.is_line({"type": "shape", "shape": "line"}) is True
    assert shape_geometry.is_line({"type": "line"}) is True
    assert shape_geometry.is_line({"type": "shape", "shape": "circle"}) is False
    assert shape_geometry.is_line(None) is False


# ============================================================================
# 2. JavaScript Shapes Line API Parity Tests
# ============================================================================

def test_js_shapes_line_api():
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js not available")

    js_code = """
    const shapes = require('./js/shapes.js');
    const assert = require('assert');

    // 1. isLine
    assert.strictEqual(shapes.isLine({ type: 'shape', shape: 'line' }), true);
    assert.strictEqual(shapes.isLine({ type: 'line' }), true);
    assert.strictEqual(shapes.isLine({ type: 'shape', shape: 'rect' }), false);

    // 2. createLine & lineGeometry
    const l1 = shapes.createLine(10, 20, 110, 70, { stroke: 'blue', strokeWidth: 4 });
    assert.strictEqual(l1.type, 'shape');
    assert.strictEqual(l1.shape, 'line');
    assert.strictEqual(l1.x, 10);
    assert.strictEqual(l1.y, 20);
    assert.strictEqual(l1.w, 100);
    assert.strictEqual(l1.h, 50);
    assert.strictEqual(l1.x1, 10);
    assert.strictEqual(l1.y1, 20);
    assert.strictEqual(l1.x2, 110);
    assert.strictEqual(l1.y2, 70);
    assert.strictEqual(l1.stroke, 'blue');
    assert.strictEqual(l1.strokeWidth, 4);

    // 3. Reversed line
    const l2 = shapes.createLine(300, 200, 100, 50);
    assert.strictEqual(l2.x, 100);
    assert.strictEqual(l2.y, 50);
    assert.strictEqual(l2.w, 200);
    assert.strictEqual(l2.h, 150);
    assert.strictEqual(l2.x1, 300);
    assert.strictEqual(l2.x2, 100);

    // 4. Zero dimensions (horizontal & vertical)
    const lh = shapes.createLine(50, 100, 250, 100);
    assert.strictEqual(lh.w, 200);
    assert.strictEqual(lh.h, 0);

    const lv = shapes.createLine(150, 40, 150, 340);
    assert.strictEqual(lv.w, 0);
    assert.strictEqual(lv.h, 300);

    // 5. Angle snapping
    const snapH = shapes.snapAngle(0, 0, 100, 5); // nearly horizontal -> snaps to 0 deg
    assert.strictEqual(snapH.y, 0);

    const snapV = shapes.snapAngle(0, 0, 5, 100); // nearly vertical -> snaps to 90 deg
    assert.strictEqual(snapV.x, 0);

    // 6. Legacy fallback
    const legacyGeom = shapes.lineGeometry({ x: 50, y: 100, w: 200, h: 20 });
    assert.strictEqual(legacyGeom.x1, 50);
    assert.strictEqual(legacyGeom.y1, 110);
    assert.strictEqual(legacyGeom.x2, 250);
    assert.strictEqual(legacyGeom.y2, 110);

    console.log('JS_LINE_TESTS_OK');
    """
    res = subprocess.run([node, "-e", js_code], capture_output=True, text=True, check=True)
    assert "JS_LINE_TESTS_OK" in res.stdout


# ============================================================================
# 3. Serialization and Deserialization Tests
# ============================================================================

def test_line_serialization_roundtrip():
    doc = PresentationDoc(
        slides=[
            SlideModel(
                id="s1",
                objects=[
                    SlideObject(
                        id="line_diag",
                        type="shape",
                        shape="line",
                        x=50,
                        y=60,
                        w=200,
                        h=150,
                        x1=50,
                        y1=60,
                        x2=250,
                        y2=210,
                        stroke="#FF0000",
                        strokeWidth=4.5,
                        opacity=0.8,
                        rotation=15.0,
                    ),
                    SlideObject(
                        id="line_horiz",
                        type="shape",
                        shape="line",
                        x=100,
                        y=300,
                        w=400,
                        h=0,
                        x1=100,
                        y1=300,
                        x2=500,
                        y2=300,
                        stroke="accent",
                    ),
                ],
            )
        ]
    )

    data = doc.model_dump(mode="json")
    json_str = json.dumps(data)
    loaded = json.loads(json_str)
    doc2 = PresentationDoc.model_validate(loaded)

    obj1 = doc2.slides[0].objects[0]
    assert obj1.id == "line_diag"
    assert obj1.x1 == 50.0
    assert obj1.y1 == 60.0
    assert obj1.x2 == 250.0
    assert obj1.y2 == 210.0
    assert obj1.stroke == "#FF0000"
    assert obj1.strokeWidth == 4.5
    assert obj1.opacity == 0.8
    assert obj1.rotation == 15.0

    obj2 = doc2.slides[0].objects[1]
    assert obj2.id == "line_horiz"
    assert obj2.w == 400.0
    assert obj2.h == 0.0
    assert obj2.x1 == 100.0
    assert obj2.x2 == 500.0


# ============================================================================
# 4. PPTX Export Tests
# ============================================================================

def test_pptx_export_native_connector():
    doc = PresentationDoc(
        slides=[
            SlideModel(
                id="s1",
                objects=[
                    # 1. Diagonal line
                    SlideObject(
                        id="l_diag",
                        type="shape",
                        shape="line",
                        x1=100,
                        y1=80,
                        x2=400,
                        y2=280,
                        stroke="#2B6CB0",
                        strokeWidth=3.0,
                    ),
                    # 2. Horizontal line (h=0)
                    SlideObject(
                        id="l_horiz",
                        type="shape",
                        shape="line",
                        x1=50,
                        y1=400,
                        x2=550,
                        y2=400,
                        stroke="#E53E3E",
                        strokeWidth=2.0,
                    ),
                    # 3. Vertical line (w=0)
                    SlideObject(
                        id="l_vert",
                        type="shape",
                        shape="line",
                        x1=600,
                        y1=100,
                        x2=600,
                        y2=500,
                        stroke="#38A169",
                        strokeWidth=2.5,
                    ),
                    # 4. Reversed line (x2 < x1, y2 < y1)
                    SlideObject(
                        id="l_rev",
                        type="shape",
                        shape="line",
                        x1=800,
                        y1=450,
                        x2=700,
                        y2=350,
                        stroke="#805AD5",
                        strokeWidth=3.0,
                    ),
                    # 5. Zero-dimension line
                    SlideObject(
                        id="l_zero",
                        type="shape",
                        shape="line",
                        x1=500,
                        y1=500,
                        x2=500,
                        y2=500,
                        stroke="#000000",
                    ),
                    # 6. Legacy line (bounding box only)
                    SlideObject(
                        id="l_legacy",
                        type="shape",
                        shape="line",
                        x=120,
                        y=600,
                        w=300,
                        h=20,
                        stroke="accent",
                    ),
                ],
            )
        ]
    )

    out = export_presentation(doc)
    prs = Presentation(out)
    slide = prs.slides[0]

    # Verify all 6 line shapes were exported
    shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(shapes) == 6

    for shp in shapes:
        # Native PowerPoint straight connector has shape_type LINE
        assert shp.shape_type == MSO_SHAPE_TYPE.LINE

    # 1. Diagonal line checks
    s0 = shapes[0]
    assert s0.begin_x == to_emu(100)
    assert s0.begin_y == to_emu(80)
    assert s0.end_x == to_emu(400)
    assert s0.end_y == to_emu(280)

    # 2. Horizontal line checks
    s1 = shapes[1]
    assert s1.begin_x == to_emu(50)
    assert s1.begin_y == to_emu(400)
    assert s1.end_x == to_emu(550)
    assert s1.end_y == to_emu(400)

    # 3. Vertical line checks
    s2 = shapes[2]
    assert s2.begin_x == to_emu(600)
    assert s2.begin_y == to_emu(100)
    assert s2.end_x == to_emu(600)
    assert s2.end_y == to_emu(500)

    # 4. Reversed line checks
    s3 = shapes[3]
    assert s3.begin_x == to_emu(800)
    assert s3.begin_y == to_emu(450)
    assert s3.end_x == to_emu(700)
    assert s3.end_y == to_emu(350)

    # 5. Zero dimension line checks
    s4 = shapes[4]
    assert s4.begin_x == to_emu(500)
    assert s4.begin_y == to_emu(500)

    # 6. Legacy line checks (horizontal across box)
    s5 = shapes[5]
    assert s5.begin_x == to_emu(120)
    assert s5.begin_y == to_emu(610)
    assert s5.end_x == to_emu(420)
    assert s5.end_y == to_emu(610)


def test_existing_rectangle_unaffected():
    # Verify standard rectangle shapes export as RECTANGLE
    doc = PresentationDoc(
        slides=[
            SlideModel(
                id="s1",
                objects=[
                    SlideObject(
                        id="rect1",
                        type="shape",
                        shape="rect",
                        x=100,
                        y=100,
                        w=200,
                        h=150,
                        fill="#E2E8F0",
                        stroke="#4A5568",
                    ),
                ],
            )
        ]
    )

    out = export_presentation(doc)
    prs = Presentation(out)
    slide = prs.slides[0]
    shp = slide.shapes[0]
    assert shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
